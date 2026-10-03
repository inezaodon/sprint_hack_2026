"""Close rules engine (Engineer 8): one test per rule on tiny hand-built harmonized + finance DBs."""
from datetime import date
from pathlib import Path

import duckdb
import pytest

from goodwill_pulse.close import enrich
from goodwill_pulse.close.rules import (CloseRun, allocate, balance_exceptions, plant_error, run_close,
                                        write_manual_workbook)
from goodwill_pulse.sources.finance import FINANCE_DDL, build as build_finance

ROOT = Path(__file__).resolve().parent.parent
RULES = ROOT / "config" / "close_rules.yaml"
SEPT = date(2026, 9, 1)


def _orders(con, rows):
    for key, ch, d, sub, ship, hand, fees, lines in rows:
        con.execute("INSERT INTO fct_orders VALUES (?,?,?,?,NULL,NULL,?,?,?,?,?,?,0,?,0,?)",
                    [key, ch, key.split(":")[1], ch, f"{d} 15:00:00+00", d, len(lines), sub, ship, hand, fees,
                     sub + ship + hand])
        for i, (sku, store, price) in enumerate(lines, 1):
            con.execute("INSERT INTO fct_order_lines VALUES (?,?,?,?,?,NULL,NULL,'x',1,?,0,?)",
                        [key, i, ch, sku, store, price, d])


@pytest.fixture
def dbs(tmp_path):
    h = tmp_path / "harmonized.duckdb"
    con = duckdb.connect(str(h))
    con.execute((ROOT / "sql" / "harmonized" / "schema.sql").read_text())
    for sku, store in [("UP-03-000001", "Store03"), ("UP-05-000002", "Store05"), ("GWM-07-000003", "Store07"),
                       ("UP-03-000005", "Store03"), ("UP-03-000006", "Store03")]:
        con.execute("INSERT INTO fct_items (item_id, store_id) VALUES (?, ?)", [sku, store])
    _orders(con, [
        ("ebay:E1", "ebay", "2026-09-05", 100, 10, 0, 13, [("UP-03-000001", "Store03", 100)]),
        ("ebay:E2", "ebay", "2026-09-20", 50, 5, 0, 6.5, [("UP-09-000004", None, 50)]),   # store filled from SKU
        ("ebay:E3", "ebay", "2026-10-01", 70, 0, 0, 7, [("UP-03-000009", "Store03", 70)]),  # October: excluded
        ("shopgoodwill:S1", "shopgoodwill", "2026-09-03", 200, 12, 3, 20, [("UP-05-000002", "Store05", 200)]),
        ("shopgoodwill:S2", "shopgoodwill", "2026-09-15", 80, 0, 0, 8, [("UP-03-000006", "Store03", 80)]),
        ("shopgoodwill:S3", "shopgoodwill", "2026-09-25", 120, 8, 3, 12, [("UP-03-000005", "Store03", 120)]),
        ("amazon:A1", "amazon", "2026-09-10", 30, 0, 0, 4.5, [("GWM-07-000010", "Store07", 30)]),
        ("goodwillfinds:G1", "goodwillfinds", "2026-09-12", 40, 0, 0, 4, [("UP-11-000011", "Store11", 40)]),
        ("goodwillbooks:GWB1", "goodwillbooks", "2026-09-08", 15, 4, 0, 2, [("GWM-07-000003", "Store07", 15)]),
    ])
    con.execute("INSERT INTO fct_refunds VALUES ('ebay:R1','ebay:E1','ebay','2026-09-07 15:00:00+00','2026-09-07',20)")
    con.execute("INSERT INTO fct_payouts VALUES ('shopgoodwill:P1','shopgoodwill','2026-09-13','2026-09-01',"
                "'2026-09-10',215,20,0,195), ('shopgoodwill:P2','shopgoodwill','2026-09-23','2026-09-11',"
                "'2026-09-20',80,8,0,72)")
    con.close()

    f = tmp_path / "finance.duckdb"
    con = duckdb.connect(str(f))
    con.execute(FINANCE_DDL)
    con.execute("""INSERT INTO bank_0101 VALUES
        ('B1','0101','2026-09-08','OSM WORLDWIDE ACH DEBIT',-150,'ACH_DEBIT','OSM-1'),
        ('B2','0101','2026-09-14','PITNEY BOWES POSTAGE BY PHONE',-60,'ACH_DEBIT','PB-1'),
        ('B3','0101','2026-09-23','EASYPOST WALLET RELOAD',-40,'ACH_DEBIT','EP-1'),
        ('B4','0101','2026-10-02','OSM WORLDWIDE ACH DEBIT',-70,'ACH_DEBIT','OSM-2'),
        ('B5','0101','2026-09-18','BNKDEPOSIT FEDEX REFUND',12.50,'BNKDEPOSIT','FXR-1'),
        ('B6','0101','2026-10-01','BNKDEPOSIT FEDEX REFUND',8.00,'BNKDEPOSIT','FXR-2'),
        ('B7','0101','2026-09-10','BNKDEPOSIT STORE CASH ECOM',99,'BNKDEPOSIT','DEP-1'),
        ('B8','0101','2026-09-11','EBAY PAYOUT',100,'ACH_CREDIT','PO-1'),
        ('B9','0101','2026-09-15','FEDEX ACH PAYMENT',-50,'ACH_DEBIT','INV-A')""")
    con.execute("""INSERT INTO fedex_invoices VALUES
        ('FX1','INV-A','2026-09-05','CHARGE','T1','2026-09-01',30,NULL,NULL),
        ('FX2','INV-A','2026-09-05','CHARGE','T2','2026-09-02',20,NULL,NULL),
        ('FX3','INV-B','2026-09-12','CHARGE','T3','2026-09-08',45,NULL,NULL),
        ('FX4','INV-OCT','2026-10-03','CHARGE','T4','2026-09-29',15,NULL,NULL),
        ('FX5','CM-INV-A','2026-09-17','REFUND','T1','2026-09-17',-12.50,'INV-A','FXR-1'),
        ('FX6','CM-INV-B','2026-09-29','REFUND','T3','2026-09-29',-8.00,'INV-B','FXR-2')""")
    con.execute("""INSERT INTO jewelry_report VALUES
        ('2026-09',1,'ebay','E1','2026-09-05','UP-03-000001','Ring',100,NULL),
        ('2026-09',2,'shopgoodwill','S1','2026-09-03','UP-05-000002','Necklace',200,'Store05'),
        ('2026-09',3,'ebay','E9','2026-09-09','UP-07-00012','Brooch',25,NULL)""")
    con.execute("""INSERT INTO gwb_statement VALUES
        ('2026-09',1,'SALE','GWB1','2026-09-08','GWM-07-000003','Item sale',15),
        ('2026-09',2,'SHIPPING','GWB1','2026-09-08',NULL,'Shipping',4),
        ('2026-09',3,'FEE','GWB1','2026-09-08',NULL,'Platform fee',-2)""")
    con.execute("INSERT INTO gwb_statement_header VALUES ('2026-09','2026-10-01',19,2,0,17,'2026-10-15')")
    con.close()
    return h, f


@pytest.fixture
def run(dbs) -> CloseRun:
    return run_close(SEPT, dbs[0], dbs[1], RULES)


def lines(run, suffix):
    a = run.allocations
    return a[a.doc_no == f"ECOM-2026-09-{suffix}"]


def amt(df, account_no, store=None, **kw):
    d = df[df.account_no == account_no]
    if store is not None:
        d = d[d.store_id == store]
    return round(d.amount.sum(), 2)


def rv(run, account_no, store=None, channel=None):
    r = run.revenue
    r = r[r.account_no == account_no]
    if store is not None:
        r = r[r.store_id == store]
    if channel is not None:
        r = r[r.channel == channel]
    return round(r.amount.sum(), 2)


def test_documents_and_columns(run):
    assert list(run.allocations.columns) == ["doc_no", "posting_date", "rule_id", "source", "account_type",
                                             "account_no", "department_code", "store_id", "description", "amount",
                                             "placeholder_account", "source_ref"]
    assert set(run.allocations.doc_no) == {f"ECOM-2026-09-{s}" for s in
                                          ["SHIPPING", "FEDEX", "SGW", "EBAY", "AMAZON", "GWF", "GWB"]}
    assert (run.allocations.posting_date == date(2026, 9, 30)).all()
    assert run.allocations.source_ref.str.contains(":").all()
    assert run.rules_version.startswith("2026.10.03-2+")
    assert all({"doc_no", "measure"} <= set(t) for t in run.source_totals)
    assert list(run.revenue.columns) == ["channel", "customer_no", "store_id", "account_no", "amount",
                                         "placeholder_account", "description", "source_ref"]


def test_v14_revenue_only_on_ar_invoice(run):
    """Close v1.4: no revenue account appears in any journal document; it is all in CloseRun.revenue."""
    rev_accounts = {"4xxxx-SGW", "4xxxx-EBAY", "4xxxx-AMZ", "4xxxx-GWF", "4xxxx-GWB", "4xxxx-SHIPINC", "4xxxx-JEWELRY"}
    assert not set(run.allocations.account_no) & rev_accounts
    assert set(run.revenue.account_no) <= rev_accounts
    assert (run.revenue.amount > 0).all() and run.revenue.placeholder_account.all()
    sales = run.revenue[~run.revenue.account_no.isin(["4xxxx-SHIPINC"])]
    assert sales.store_id.notna().all()                    # store credit on every sales line
    bal = run.customer_balances().set_index("channel")
    assert bal.loc["ebay"].to_dict() == {"customer_no": "CUST-EBAY", "invoice_total": 165.0, "fees": 19.5,
                                         "refunds": 20.0, "payouts": 100.0, "open_balance": 25.5}
    assert round(run.revenue.amount.sum(), 2) == 100 + 50 + 15 + 215 + 80 + 131 + 30 + 40 + 19


def test_every_document_balances(run):
    t = run.doc_totals()
    assert (t.balance == 0).all(), t
    assert not [e for e in run.exceptions if e["rule_id"] == "balance_check"]


def test_unbalanced_document_is_an_exception(run):
    bad = run.allocations.copy()
    bad.loc[bad.index[0], "amount"] += 0.01
    exc = balance_exceptions(bad)
    assert len(exc) == 1 and exc[0]["severity"] == "error" and exc[0]["amount"] == 0.01


def test_rule_shipping_postage_bank_0101_to_gl_10009(run):
    s = lines(run, "SHIPPING")
    dr = s[s.account_no == "10009"]
    assert sorted(dr.amount) == [40.0, 60.0, 150.0]        # October OSM line excluded, non-postage lines ignored
    assert set(dr.source_ref) == {"bank_0101:B1", "bank_0101:B2", "bank_0101:B3"}
    cr = s[(s.account_type == "Bank Account") & (s.account_no == "0101")]
    assert round(cr.amount.sum(), 2) == -250.0 and len(cr) == 3


def test_rule_fedex_gl_40356_dept_180_vendor_v00122_net_of_bnkdeposit(run):
    f = lines(run, "FEDEX")
    exp = f[f.account_no == "40356"]
    assert (exp.department_code == "180").all()
    assert sorted(exp.amount) == [-12.5, 45.0, 50.0]        # INV-A 50, INV-B 45, Sept BNKDEPOSIT refund -12.50
    v = f[(f.account_type == "Vendor") & (f.account_no == "V00122")]
    assert len(v) == 1 and v.amount.iloc[0] == -82.5
    tot = next(t for t in run.source_totals if t["doc_no"] == "ECOM-2026-09-FEDEX")
    assert tot["file_total"] == 74.5 and tot["rules_total"] == 82.5     # difference = cross-month refund 8.00


def test_cross_month_fedex_refund_is_an_exception_not_netted(run):
    f = lines(run, "FEDEX")
    assert "bank_0101:B6" not in set(f.source_ref)
    exc = [e for e in run.exceptions if "Cross-month" in e["message"]]
    assert len(exc) == 1 and exc[0]["amount"] == 8.0 and exc[0]["severity"] == "warning"
    assert "2026-10-01" in exc[0]["message"] and exc[0]["owner"]


def test_rule_shopgoodwill_period1_periodic_only_period3_all_reports(run):
    s = lines(run, "SGW")
    p1 = s[s.description.str.startswith("P1 ")]
    assert set(p1.source_ref.str.split("/").str[0]) == {"harmonized.fct_payouts:shopgoodwill:P1"}
    assert amt(p1, "6xxxx-FEES") == 20 and amt(p1, "CUST-SGW") == -20
    r = run.revenue[run.revenue.channel == "shopgoodwill"].set_index("account_no")
    ship_ref = r.loc["4xxxx-SHIPINC", "source_ref"]           # P1 from the statement, P3 from detail reports
    assert "fct_payouts:shopgoodwill:P1/shipping" in ship_ref and "fct_orders:shopgoodwill/2026-09-21" in ship_ref
    assert round(r.amount.sum(), 2) == 215 + 80 + 131          # statement P1 + statement P2 + detail P3 (120 + 11)
    p3 = s[s.description.str.startswith("P3 ")]
    assert p3.source_ref.str.startswith("harmonized.fct_").all() and not p3.source_ref.str.contains("payouts").any()
    assert amt(p3, "6xxxx-FEES") == 12
    assert rv(run, "4xxxx-SGW", "Store03") == 200              # P2 statement 80 + P3 detail 120
    assert rv(run, "4xxxx-SHIPINC", channel="shopgoodwill") == 26  # P1 statement share 15 + P3 detail 11


def test_shopgoodwill_missing_period_statement_falls_back_with_exception(dbs):
    con = duckdb.connect(str(dbs[0]))
    con.execute("DELETE FROM fct_payouts WHERE payout_key = 'shopgoodwill:P1'")
    con.close()
    r = run_close(SEPT, dbs[0], dbs[1], RULES)
    assert any(e["severity"] == "error" and "Period 1" in e["message"] for e in r.exceptions)
    p1 = lines(r, "SGW")
    p1 = p1[p1.description.str.startswith("P1 ")]
    assert p1.source_ref.str.startswith("harmonized.fct_orders").all() and amt(p1, "6xxxx-FEES") == 20
    assert rv(r, "4xxxx-SHIPINC", channel="shopgoodwill") == 26


def test_rule_channel_revenue_ebay_store_credit_fees_refunds_payouts(run):
    e = lines(run, "EBAY")
    assert amt(e, "6xxxx-FEES") == 19.5 and amt(e, "4xxxx-REFUNDS", "Store03") == 20
    bank = e[e.account_type == "Bank Account"]
    assert list(bank.amount) == [100.0] and list(bank.source_ref) == ["bank_0101:B8"]
    assert amt(e, "CUST-EBAY") == -139.5
    assert rv(run, "4xxxx-EBAY", "Store09") == 50              # NULL store filled from SKU prefix (enrich)
    assert rv(run, "4xxxx-SHIPINC", channel="ebay") == 15
    assert rv(run, "4xxxx-AMZ", "Store07") == 30 and rv(run, "4xxxx-GWF", "Store11") == 40
    assert amt(lines(run, "GWF"), "CUST-GWF") == -4


def test_rule_goodwillbooks_statement(run):
    g = lines(run, "GWB")
    assert amt(g, "6xxxx-FEES") == 2 and amt(g, "CUST-GWB") == -2
    assert g.source_ref.str.startswith("gwb_statement:").all()
    assert rv(run, "4xxxx-GWB", "Store07") == 15 and rv(run, "4xxxx-SHIPINC", channel="goodwillbooks") == 4
    tot = next(t for t in run.source_totals if t["doc_no"] == "ECOM-2026-09-GWB")
    assert tot["file_total"] == tot["loaded_total"] == tot["rules_total"] == 19


def test_goodwillbooks_statement_lines_must_sum_to_total(dbs):
    con = duckdb.connect(str(dbs[1]))
    con.execute("UPDATE gwb_statement_header SET net = 18")
    con.close()
    r = run_close(SEPT, dbs[0], dbs[1], RULES)
    assert any(e["rule_id"] == "goodwillbooks_statement" and e["severity"] == "error" for e in r.exceptions)


def test_rule_jewelry_copivot_fill_and_unmatched_exception(run):
    assert rv(run, "4xxxx-JEWELRY", "Store03", "ebay") == 100   # Supplier blank -> filled by SKU lookup (Co-Pivot)
    assert rv(run, "4xxxx-JEWELRY", "Store05", "shopgoodwill") == 200
    assert rv(run, "4xxxx-EBAY", "Store03") == 0 and rv(run, "4xxxx-SGW", "Store05") == 0   # moved, not doubled
    exc = [e for e in run.exceptions if e["rule_id"] == "jewelry_reclass"]
    assert len(exc) == 1 and exc[0]["severity"] == "error" and exc[0]["amount"] == 25.0
    assert "UP-07-00012" in exc[0]["message"]
    tot = next(t for t in run.source_totals if t["source"] == "jewelry_report")
    assert tot["doc_no"] is None and tot["file_total"] == 325 and tot["rules_total"] == 300


def test_enrich_fill_supplier():
    import pandas as pd
    df = pd.DataFrame({"sku": ["A", "B", "C"], "supplier": ["Store01", None, " "]})
    out = enrich.fill_supplier(df, {"B": "Store02", "A": "Store09"})
    assert list(out.store_id) == ["Store01", "Store02", None]
    assert list(out.supplier_source) == ["report", "co-pivot", "unmatched"]
    assert enrich.store_from_sku("gwm-07-123456") == "Store07" and enrich.store_from_sku("UP-07-00012") is None


def test_missing_sources_raise_exceptions(dbs, tmp_path):
    r = run_close(SEPT, dbs[0], tmp_path / "nope.duckdb", RULES)
    msgs = [e["message"] for e in r.exceptions if e["severity"] == "error"]
    assert any("finance DB not found" in m for m in msgs)
    assert sum(m.startswith("Missing source") for m in msgs) >= 4      # shipping, FedEx, GWB statement, jewelry
    assert (r.doc_totals().balance == 0).all()
    assert "ECOM-2026-09-GWB" in set(r.allocations.doc_no)           # accrued from detail
    r2 = run_close(date(2026, 3, 1), dbs[0], dbs[1], RULES)          # a month with no data at all
    assert r2.allocations.empty and len([e for e in r2.exceptions if e["severity"] == "error"]) >= 7


def test_allocate_sums_exactly():
    out = allocate(1000, {"a": 1, "b": 1, "c": 1})
    assert sum(out.values()) == 1000 and sorted(out.values()) == [333, 333, 334]


def test_manual_workbook_has_exactly_one_planted_error(run, tmp_path):
    from openpyxl import load_workbook
    path = write_manual_workbook(run, tmp_path, seed=7)
    assert path.name == "E-Commerce Allocation 2026-09.xlsx" and (tmp_path / "ANSWER_KEY.md").exists()
    wb = load_workbook(path)
    assert {"Journal Entries", "Inputs", "Invoices"} <= set(wb.sheetnames)
    assert wb["Invoices"].max_row == len(run.revenue) + 1
    rows = list(wb["Journal Entries"].iter_rows(values_only=True))
    assert list(rows[0]) == ["doc_no", "posting_date", "source", "account_type", "account_no", "department_code",
                             "store_id", "description", "amount"]
    got = [r[8] for r in rows[1:]]
    want = list(run.allocations.amount)
    diffs = [i for i, (a, b) in enumerate(zip(got, want)) if round(a - b, 2) != 0]
    assert len(got) == len(want) and len(diffs) == 2                  # the keyed line + its balancing line
    docs = {run.allocations.doc_no.iloc[i] for i in diffs}
    assert len(docs) == 1
    orange = [c for row in wb["Inputs"].iter_rows() for c in row if c.fill.fgColor.rgb in ("00FFC000", "FFFFC000")]
    assert orange
    key = (tmp_path / "ANSWER_KEY.md").read_text()
    assert docs.pop() in key
    _, err = plant_error(run.allocations, 7)
    assert sorted(str(int(abs(err["correct"]))))== sorted(str(int(abs(err["keyed"]))))   # digits transposed


# --------------------------------------------------------------------------------------- finance source DB build
def _tiny_truth(path):
    con = duckdb.connect(str(path))
    con.execute("""
        CREATE TABLE items (item_id VARCHAR, store_id VARCHAR, category VARCHAR, title VARCHAR);
        CREATE TABLE orders (order_id VARCHAR, channel VARCHAR, marketplace_order_id VARCHAR, paid_at TIMESTAMPTZ,
            shipping_charged DECIMAL(12,2), handling DECIMAL(12,2), marketplace_fee DECIMAL(12,2),
            payment_fee DECIMAL(12,2));
        CREATE TABLE order_lines (order_id VARCHAR, line_no INTEGER, item_id VARCHAR, sale_price DECIMAL(12,2));
        CREATE TABLE refunds (refund_id VARCHAR, order_id VARCHAR, refunded_at TIMESTAMPTZ, amount DECIMAL(12,2));
        CREATE TABLE payouts (payout_id VARCHAR, channel VARCHAR, period_start DATE, period_end DATE, paid_on DATE,
            gross DECIMAL(12,2), fees DECIMAL(12,2), refunds DECIMAL(12,2), net DECIMAL(12,2));
        CREATE TABLE shipping_charges (charge_id VARCHAR, carrier VARCHAR, order_id VARCHAR, charged_at TIMESTAMPTZ,
            amount DECIMAL(12,2), tracking VARCHAR);
    """)
    for i in range(1, 7):
        con.execute("INSERT INTO items VALUES (?,?,?,?)", [f"UP-0{i}-00000{i}", f"Store0{i}", "Jewelry", f"Ring {i}"])
        con.execute("INSERT INTO orders VALUES (?,?,?,?,5,3,4,1)",
                    [f"O{i}", "shopgoodwill", f"{9000 + i}", f"2026-09-{i * 4:02d} 16:00:00+00"])
        con.execute("INSERT INTO order_lines VALUES (?,1,?,?)", [f"O{i}", f"UP-0{i}-00000{i}", 50 + i])
        con.execute("INSERT INTO shipping_charges VALUES (?,?,?,?,?,?)",
                    [f"C{i}", "fedex", f"O{i}", f"2026-09-{i * 4:02d} 18:00:00+00", 10 + i, f"TRK{i}"])
        con.execute("INSERT INTO shipping_charges VALUES (?,?,?,?,?,?)",
                    [f"P{i}", "osm", f"O{i}", f"2026-09-{i * 4:02d} 18:00:00+00", 4 + i, f"OSM{i}"])
    con.execute("INSERT INTO items VALUES ('GWM-03-000100','Store03','Books','A book')")
    con.execute("INSERT INTO orders VALUES ('B1','goodwillbooks','GWB000001','2026-09-09 16:00:00+00',4,0,2,0)")
    con.execute("INSERT INTO order_lines VALUES ('B1',1,'GWM-03-000100',12)")
    con.execute("INSERT INTO payouts VALUES ('PO1','shopgoodwill','2026-09-01','2026-09-10','2026-09-14',100,10,0,90)")
    con.close()


def test_finance_build_plants_cross_month_refund_and_unmatched_jewelry(tmp_path):
    truth = tmp_path / "_truth.duckdb"
    _tiny_truth(truth)
    out = tmp_path / "finance.duckdb"
    counts = build_finance(truth, out, seed=7)
    assert counts["bank_0101"] > 0 and counts["fedex_invoices"] >= 7 and counts["jewelry_report"] == 6
    con = duckdb.connect(str(out), read_only=True)
    kinds = {k for (k,) in con.execute("SELECT kind FROM _dirty_data").fetchall()}
    assert kinds == {"cross_month_refund", "unmatched_sku"}
    dep = con.execute("SELECT post_date, amount FROM bank_0101 WHERE type='BNKDEPOSIT' "
                      "AND description LIKE '%FEDEX%'").fetchall()
    assert [d for d, _ in dep] == [date(2026, 10, 1)]
    cm = con.execute("SELECT invoice_date, orig_invoice_no FROM fedex_invoices WHERE line_type='REFUND'").fetchall()
    inv_date = con.execute("SELECT DISTINCT invoice_date FROM fedex_invoices WHERE invoice_no = ?",
                           [cm[0][1]]).fetchone()[0]
    assert cm[0][0] == date(2026, 9, 29) and inv_date.month == 9
    assert con.execute("SELECT count(*) FROM jewelry_report WHERE length(sku) = 11 AND supplier IS NULL"
                       ).fetchone()[0] == 1
    assert con.execute("SELECT count(*) FROM bank_0101 WHERE description LIKE 'OSM WORLDWIDE%'").fetchone()[0] >= 1
    h = con.execute("SELECT net FROM gwb_statement_header WHERE statement_month='2026-09'").fetchone()[0]
    s = con.execute("SELECT sum(amount) FROM gwb_statement WHERE statement_month='2026-09'").fetchone()[0]
    assert h == s == 14
    con.close()
    out2 = tmp_path / "finance2.duckdb"
    build_finance(truth, out2, seed=7)
    a, b = duckdb.connect(str(out), read_only=True), duckdb.connect(str(out2), read_only=True)
    for t in ["bank_0101", "fedex_invoices", "jewelry_report", "gwb_statement"]:
        assert a.execute(f"SELECT * FROM {t} ORDER BY ALL").fetchall() == b.execute(f"SELECT * FROM {t} ORDER BY ALL").fetchall()
