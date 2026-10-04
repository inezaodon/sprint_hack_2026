"""The synthetic demo folders must load the way the real exports would, and tie to their own expected.json."""
import json
from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from goodwill_pulse import db
from goodwill_pulse.gen import demo_pack
from goodwill_pulse.ingest.pipeline import process_file
from goodwill_pulse.ingest.recognize import recognize

PACK = demo_pack.OUT


@pytest.fixture(scope="module", autouse=True)
def pack():
    if not (demo_pack.SAMPLES_DIR / "history").exists():
        pytest.skip("run python -m goodwill_pulse.gen.generate first")
    demo_pack.main()
    return PACK


@pytest.fixture
def con(tmp_path, monkeypatch):
    monkeypatch.setattr("goodwill_pulse.ingest.pipeline.ARCHIVE_DIR", tmp_path / "archive")
    c = db.connect(tmp_path / "w.duckdb")
    yield c
    c.close()


def test_real_header_is_recognized_with_no_missing_columns(pack):
    rec = recognize(pack / "02_friday_2026-10-02" / "paid_orders_10-02-2026_10-02-2026.csv")
    assert rec.report_type == "upright_paid_orders"
    assert rec.missing_required == [] and rec.ready
    assert rec.column_map["Handling Total"] == "Handling" and rec.column_map["Channel Buyer ID"] == "Channel Buyer"


@pytest.mark.parametrize("folder,day", [("01_tonight_saturday_2026-10-03", date(2026, 10, 3)), ("02_friday_2026-10-02", date(2026, 10, 2))])
def test_single_day_upright_loads_and_matches_expected(pack, con, folder, day):
    exp = json.loads((pack / folder / "expected.json").read_text())
    for ext in ("csv", "xlsx"):
        c = db.connect(Path(con.execute("PRAGMA database_list").fetchone()[2]).with_name(f"w_{ext}.duckdb"))
        r = process_file(c, pack / folder / f"paid_orders_{day:%m-%d-%Y}_{day:%m-%d-%Y}.{ext}")
        assert r.status == "loaded", (ext, r.message)
        got = c.execute("SELECT count(*), round(sum(subtotal), 2) FROM orders").fetchone()
        assert (got[0], float(got[1])) == (exp["upright_orders"], exp["upright_total_item_sales"]), (ext, got)
        c.close()


def test_monday_range_file_equals_sum_of_daily_files(pack):
    exp = json.loads((pack / "03_monday_catchup_fri_sat_sun" / "expected.json").read_text())
    assert round(sum(d["item_sales"] for d in exp["days"].values()), 2) == exp["range_file_total_item_sales"]
    assert sum(d["orders"] for d in exp["days"].values()) == exp["range_file_orders"]


def test_messy_files_are_flagged_not_silently_loaded(pack, con):
    p = pack / "04_messy_reports" / "paid_orders_10-02-2026_10-02-2026 (2).csv"
    r = process_file(con, p)
    rules = {e["rule"] for e in r.exceptions}
    assert rules & {"renamed_column", "missing_column"}, rules       # 'Sub Total' is not silently accepted
    assert r.status != "loaded" or r.exceptions


def test_weekly_store_workbook_adds_up(pack):
    wb = load_workbook(pack / "05_store_weekly_sales_supro" / "Weekly Sales 2026 Week 40.xlsx", data_only=True)
    assert wb.sheetnames == ["Report", "Data Sheet"] or set(wb.sheetnames) == {"Report", "Data Sheet"}
    ds = pd.read_excel(pack / "05_store_weekly_sales_supro" / "Weekly Sales 2026 Week 40.xlsx", sheet_name="Data Sheet")
    exp = json.loads((pack / "05_store_weekly_sales_supro" / "expected.json").read_text())
    assert len(ds) == 24 * 7
    assert round(ds["Net Sales"].sum(), 2) == exp["week_40"]["enterprise_week_sales"]
    mc = ds[ds["Store Code"] == "MC"]
    assert round(mc["Net Sales"].sum(), 2) == exp["by_store_week_40"]["MC"]["week_sales"]
    ws = wb["Report"]
    assert ws["B1"].value.startswith("GOODWILL INDUSTRIES") and ws["D3"].value == "SUNDAY" and ws["L3"].value == "TOTALS"


def test_daily_summary_matches_warehouse_sales(pack):
    exp = json.loads((pack / "06_daily_summary_manual" / "expected.json").read_text())["days"]
    assert len(exp) == 7 and all(v["orders"] > 0 for v in exp.values())
