"""Close outputs (Engineer 9): journal, BC Excel, AR invoice, reconciliation, workbook compare, /api/close."""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from goodwill_pulse.close import bc_excel, invoice, reconcile, workbook_compare
from goodwill_pulse.close.journal import BC_COLUMNS, UnbalancedJournalError, build_journal

SEPT = date(2026, 9, 1)
PD = date(2026, 9, 30)


from goodwill_pulse.close.rules import REVENUE_COLUMNS  # noqa: E402


@dataclass
class FakeRun:
    month: date
    rules_version: str
    allocations: pd.DataFrame
    exceptions: list = field(default_factory=list)
    source_totals: list = field(default_factory=list)
    # v1.4: per-channel revenue booked by the AR invoice (empty here = pre-v1.4 style run)
    revenue: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=REVENUE_COLUMNS))


def _row(doc, source, atype, acct, amount, desc, dept="", store="", rule="r", ph=False, ref="t:1"):
    return {"doc_no": doc, "posting_date": PD, "rule_id": rule, "source": source, "account_type": atype,
            "account_no": acct, "department_code": dept, "store_id": store, "description": desc,
            "amount": amount, "placeholder_account": ph, "source_ref": ref}


def fake_run(unbalanced: bool = False, exceptions=None) -> FakeRun:
    rows = [
        _row("ECOM-2026-09-FEDEX", "fedex_invoices", "G/L Account", "40356", 1234.56, "FedEx charges Sept",
             dept="180", rule="fedex_net_refunds"),
        _row("ECOM-2026-09-FEDEX", "fedex_invoices", "G/L Account", "40356", -42.10, "FedEx refunds (BNKDEPOSIT)",
             dept="180", rule="fedex_net_refunds"),
        _row("ECOM-2026-09-FEDEX", "fedex_invoices", "Vendor", "V00122",
             -1192.46 + (5 if unbalanced else 0), "FedEx net payable", rule="fedex_net_refunds"),
        _row("ECOM-2026-09-SHIPPING", "bank_0101", "G/L Account", "10009", 800.00, "OSM Worldwide postage"),
        _row("ECOM-2026-09-SHIPPING", "bank_0101", "Bank Account", "0101", -800.00, "OSM Worldwide postage"),
        _row("ECOM-2026-09-EBAY", "ebay", "Customer", "CUST-EBAY", 900.00, "eBay receivable", ph=True),
        _row("ECOM-2026-09-EBAY", "ebay", "G/L Account", "6xxxx-FEES", 100.00, "eBay fees", ph=True),
        _row("ECOM-2026-09-EBAY", "ebay", "G/L Account", "4xxxx-EBAY", -600.00, "eBay revenue Store01",
             store="Store01", ph=True),
        _row("ECOM-2026-09-EBAY", "ebay", "G/L Account", "4xxxx-EBAY", -400.00, "eBay revenue Store02",
             store="Store02", ph=True),
    ]
    totals = [
        {"source": "fedex_invoices", "doc_no": "ECOM-2026-09-FEDEX", "measure": "FedEx charges net of refunds",
         "rows": 30, "file_total": 1192.46, "loaded_total": 1192.46, "rules_total": 1192.46},
        {"source": "bank_0101", "doc_no": "ECOM-2026-09-SHIPPING", "measure": "postage paid from bank 0101",
         "rows": 12, "file_total": 800.00, "loaded_total": 800.00, "rules_total": 800.00},
        {"source": "harmonized.fct_orders:ebay", "doc_no": "ECOM-2026-09-EBAY",
         "measure": "item revenue (order subtotal)", "rows": 50, "file_total": 1000.00, "loaded_total": 1000.00,
         "rules_total": 1000.00},
    ]
    return FakeRun(SEPT, "test-1", pd.DataFrame(rows), list(exceptions or []), totals)


# --- journal ------------------------------------------------------------------------------------------------------
def test_balanced_documents():
    j = build_journal(fake_run())
    assert [d.document_no for d in j.documents] == ["ECOM-2026-09-EBAY", "ECOM-2026-09-FEDEX",
                                                     "ECOM-2026-09-SHIPPING"]
    assert j.balanced and all(d.balanced for d in j.documents)
    fedex = j.document("ECOM-2026-09-FEDEX")
    assert fedex.total_debit == Decimal("1234.56") and fedex.total_credit == Decimal("1234.56")
    assert [l.line_no for l in fedex.lines] == [10000, 20000, 30000]
    assert fedex.lines[0].department_code == "180"
    assert not [i for i in j.issues if i["severity"] == "error"]
    assert any(i["rule_id"] == "placeholder_accounts" for i in j.issues)


def test_unbalanced_document_flagged_and_strict_raises():
    j = build_journal(fake_run(unbalanced=True))
    assert not j.balanced
    errs = [i for i in j.issues if i["rule_id"] == "journal_balance"]
    assert len(errs) == 1 and "ECOM-2026-09-FEDEX" in errs[0]["message"] and errs[0]["severity"] == "error"
    with pytest.raises(UnbalancedJournalError):
        build_journal(fake_run(unbalanced=True), strict=True)


# --- BC Excel -----------------------------------------------------------------------------------------------------
def test_excel_round_trip(tmp_path):
    j = build_journal(fake_run())
    path = bc_excel.write_journal_xlsx(j, tmp_path)
    wb = load_workbook(path)
    assert wb.sheetnames[:2] == ["All lines", "Control totals"]
    assert set(wb.sheetnames[2:]) == {d.document_no for d in j.documents}
    ws = wb["All lines"]
    assert [c.value for c in ws[1]] == BC_COLUMNS
    assert ws.freeze_panes == "A2"
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(j.lines)
    di, ci = BC_COLUMNS.index("Debit Amount"), BC_COLUMNS.index("Credit Amount")
    assert round(sum(r[di] or 0 for r in rows), 2) == round(sum(r[ci] or 0 for r in rows), 2) == 3034.56
    assert rows[0][BC_COLUMNS.index("Posting Date")].date() == PD          # typed date cell
    ct = wb["Control totals"]
    total = next(r for r in ct.iter_rows(values_only=True) if r[0] == "TOTAL")
    assert total[4] == total[5] == 3034.56 and total[7] == "Yes"
    # deterministic: same journal -> same bytes
    p2 = bc_excel.write_journal_xlsx(j, tmp_path / "again")
    assert path.read_bytes() == p2.read_bytes()


# --- AR invoice ---------------------------------------------------------------------------------------------------
def test_invoice_json_shape(tmp_path):
    j = build_journal(fake_run())
    invs, issues = invoice.build_invoices(j)
    assert not issues
    assert len(invs) == 1 and invs[0].customer_number == "CUST-EBAY" and invs[0].placeholder
    assert invs[0].total == Decimal("900.00") and invs[0].ties_out
    doc = json.loads(invoice.write_invoice_json(j, invs, tmp_path).read_text())
    assert "CUST-EBAY" in doc["_placeholders"]
    p = doc["salesInvoices"][0]
    assert {"customerNumber", "invoiceDate", "postingDate", "salesInvoiceLines"} <= set(p)
    assert p["invoiceDate"] == "2026-09-30"
    line = p["salesInvoiceLines"][0]
    assert {"lineType", "lineObjectNumber", "description", "quantity", "unitPrice"} <= set(line)
    assert line["lineType"] == "Account"
    assert round(sum(l["quantity"] * l["unitPrice"] for l in p["salesInvoiceLines"]), 2) == 900.00
    xl = load_workbook(invoice.write_invoice_xlsx(j, invs, tmp_path))
    assert xl["Invoices"]["C2"].value == "PLACEHOLDER" and xl["Invoices"]["H2"].value == 900.0


# --- reconciliation + approval ------------------------------------------------------------------------------------
def test_reconciliation_diffs_and_status():
    run = fake_run()
    run.source_totals[0]["loaded_total"] = 1158.00                       # 34.46 short, unexplained
    run.source_totals.append({"source": "jewelry_report", "rows": 0, "file_total": None, "loaded_total": None,
                              "rules_total": None})
    rows = {r.source: r for r in reconcile.reconcile(run.source_totals, build_journal(run))}
    assert rows["bank_0101"].status == "ok" and rows["bank_0101"].diff_rules_journal == 0
    eb = rows["harmonized.fct_orders:ebay"]
    assert eb.status == "ok" and eb.journal_total == Decimal("1000.00")          # revenue credits, not debits
    f = rows["fedex_invoices"]
    assert f.journal_total == Decimal("1192.46")                                # vendor net, not gross debits
    assert f.status == "break" and f.diff_file_loaded == Decimal("34.46")
    assert rows["jewelry_report"].status == "missing"
    # explained by an exception with the same amount
    exc = [{"source": "fedex_invoices", "rule_id": "x", "severity": "warning", "message": "late file",
            "amount": 34.46}]
    rows2 = {r.source: r for r in reconcile.reconcile(run.source_totals, build_journal(run), exc)}
    assert rows2["fedex_invoices"].status == "explained"
    # explained by several exceptions that sum to the difference (e.g. held-out jewelry SKUs)
    exc3 = [dict(exc[0], amount=30.00), dict(exc[0], amount=4.46, message="another")]
    rows3 = {r.source: r for r in reconcile.reconcile(run.source_totals, build_journal(run), exc3)}
    assert rows3["fedex_invoices"].status == "explained" and "2 exceptions" in rows3["fedex_invoices"].note


def test_approval_gate():
    err = {"source": "fedex_invoices", "rule_id": "x", "severity": "error", "message": "boom", "amount": 1}
    a = reconcile.approve([err])
    assert a.blocked and not a.approved
    assert not reconcile.approve([err], override=True, note="  ").approved          # override needs a note
    ok = reconcile.approve([err], override=True, note="Controller accepts", approver="Amanda")
    assert ok.approved and ok.override
    assert reconcile.approve([err | {"severity": "warning"}]).approved
    assert reconcile.approve([err], resolved_ids={reconcile.exception_id(err)}).approved


# --- workbook compare ---------------------------------------------------------------------------------------------
def _write_workbook(path, run, mutate=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal Entries"
    cols = ["doc_no", "posting_date", "source", "account_type", "account_no", "department_code", "store_id",
            "description", "amount"]
    ws.append(["E-Commerce Allocation 2026-09"])                           # title row above header
    ws.append(cols)
    for i, r in enumerate(run.allocations.to_dict("records")):
        vals = [r[c] for c in cols]
        if mutate:
            vals = mutate(i, vals)
        ws.append(vals)
    wb.save(path)
    return path


def test_workbook_compare_finds_planted_transposition(tmp_path):
    run = fake_run()
    def plant(i, vals):
        if i == 0:
            vals[-1] = 1243.56                                              # 1234.56 re-keyed as 1243.56
        return vals
    path = _write_workbook(tmp_path / "wb.xlsx", run, plant)
    res = workbook_compare.compare(run, path)                               # PM's e2e signature
    assert res["matches"] == len(run.allocations) - 1
    assert len(res["differences"]) == 1
    d = res["differences"][0]
    assert d["kind"] == "amount" and d["doc_no"] == "ECOM-2026-09-FEDEX" and d["account_no"] == "40356"
    assert d["system_amount"] == 1234.56 and d["workbook_amount"] == 1243.56 and d["difference"] == 9.0
    assert "Transposed" in d["explanation"] and "out of balance by 9.00" in d["explanation"]
    assert d["workbook_row"] == 3


def test_workbook_compare_catches_rules_planted_error(tmp_path):
    """Engineer 8's planted error moves an input AND its balancing line; we must report it once."""
    rules = pytest.importorskip("goodwill_pulse.close.rules")
    run = fake_run()
    path = rules.write_manual_workbook(run, tmp_path)
    res = workbook_compare.compare(run, path)
    assert len(res["differences"]) == 1, res["differences"]
    d = res["differences"][0]
    assert d["doc_no"] == "ECOM-2026-09-FEDEX" and d["account_no"] == "40356" and d["system_amount"] == 1234.56
    assert d["follow_on"]["account_no"] == "V00122"
    assert "Transposed" in d["explanation"] and "still balances" in d["explanation"]
    assert res["matches"] == len(run.allocations) - 2


def test_workbook_compare_clean_and_coding_diff(tmp_path):
    run = fake_run()
    assert workbook_compare.compare(run, _write_workbook(tmp_path / "a.xlsx", run))["status"] == "match"
    def recode(i, vals):
        if i == 7:
            vals[6] = "Store03"
        return vals
    res = workbook_compare.compare(run, _write_workbook(tmp_path / "b.xlsx", run, recode))
    assert [d["kind"] for d in res["differences"]] == ["coding"]
    assert "Store01 → Store03" in res["differences"][0]["explanation"]
    assert workbook_compare.compare(run, tmp_path / "nope.xlsx")["status"] == "no_workbook"


# --- API ----------------------------------------------------------------------------------------------------------
@pytest.fixture
def client(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from goodwill_pulse.routes import close as rc

    state = {"run": fake_run()}
    monkeypatch.setattr(rc, "_run_close", lambda month: state["run"])
    monkeypatch.setattr(rc, "OUT_ROOT", tmp_path / "out")
    monkeypatch.setattr(rc, "CLOSE_INPUTS_DIR", tmp_path / "inputs")
    monkeypatch.setattr(rc, "HARMONIZED_PATH", tmp_path / "missing.duckdb")
    monkeypatch.setattr(rc, "_cache", {})
    (tmp_path / "inputs" / "2026-09").mkdir(parents=True)
    _write_workbook(workbook_compare.workbook_path(tmp_path / "inputs", SEPT), state["run"])
    app = FastAPI()
    app.include_router(rc.router)
    c = TestClient(app)
    c.state = state
    return c


def test_api_run_approve_export(client, tmp_path):
    r = client.get("/api/close/run?month=2026-09")
    assert r.status_code == 200, r.text
    body = r.json()
    assert [s["label"] for s in body["stages"]] == ["Acquire", "Archive", "Enrich", "Apply rules", "BC output",
                                                    "Post + reconcile"]
    assert len(body["sources"]) == 9
    assert body["journal"]["balanced"] and len(body["journal"]["documents"]) == 3
    assert body["comparison"]["status"] == "match"
    assert body["can_approve"]
    assert client.get("/api/close/export?month=2026-09&kind=journal").status_code == 409   # not approved yet
    rec = client.get("/api/close/export?month=2026-09&kind=reconciliation")
    assert rec.status_code == 200 and rec.content[:2] == b"PK"
    a = client.post("/api/close/approve?month=2026-09", json={"approver": "Amanda"})
    assert a.status_code == 200 and a.json()["approval"]["approved"]
    for kind in ("journal", "invoice", "invoice_json"):
        assert client.get(f"/api/close/export?month=2026-09&kind={kind}").status_code == 200
    assert (tmp_path / "out" / "2026-09" / "close_run.json").exists()
    assert client.get("/api/close/export?month=2026-09&kind=bogus").status_code == 422


def test_api_approve_blocked_then_override(client):
    client.state["run"] = fake_run(exceptions=[{"source": "fedex_invoices", "rule_id": "fedex_late_refund",
                                                "severity": "error", "message": "Refund dated Oct 1",
                                                "amount": 42.10, "owner": "AP"}])
    body = client.get("/api/close/run?month=2026-09").json()
    assert not body["can_approve"] and body["blocking_count"] == 1
    assert client.post("/api/close/approve?month=2026-09", json={}).status_code == 409
    ok = client.post("/api/close/approve?month=2026-09", json={"override": True, "note": "Accrued manually"})
    assert ok.status_code == 200 and ok.json()["approval"]["override"]


def test_api_503_when_not_ready(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from goodwill_pulse.routes import close as rc

    def boom(month):
        raise rc.CloseNotReady("The close rules engine is not available yet")
    monkeypatch.setattr(rc, "_run_close", boom)
    monkeypatch.setattr(rc, "_cache", {})
    monkeypatch.setattr(rc, "HARMONIZED_PATH", tmp_path / "missing.duckdb")
    app = FastAPI()
    app.include_router(rc.router)
    c = TestClient(app)
    r = c.get("/api/close/run?month=2026-09")
    assert r.status_code == 503 and "not available" in r.json()["detail"]
    assert c.get("/api/close/months").status_code == 503


# --- Close v1.4: revenue only on the AR invoice --------------------------------------------------------------------
def fake_run_v14(revenue_left_in_journal: bool = False) -> FakeRun:
    """eBay doc carries only fees, refunds and the payout against CUST-EBAY; revenue comes in a per-channel table."""
    run = fake_run()
    a = run.allocations
    a = a[a["doc_no"] != "ECOM-2026-09-EBAY"]
    ebay = [
        _row("ECOM-2026-09-EBAY", "ebay", "G/L Account", "6xxxx-FEES", 100.00, "eBay fees", ph=True),
        _row("ECOM-2026-09-EBAY", "ebay", "G/L Account", "4xxxx-REFUNDS", 50.00, "eBay refunds", ph=True),
        _row("ECOM-2026-09-EBAY", "ebay", "Bank Account", "0101", 700.00, "eBay payouts to 1st Source 0101"),
        _row("ECOM-2026-09-EBAY", "ebay", "Customer", "CUST-EBAY", -850.00, "eBay settled", ph=True),
    ]
    if revenue_left_in_journal:
        ebay += [_row("ECOM-2026-09-EBAY", "ebay", "G/L Account", "4xxxx-EBAY", -25.00, "stray revenue", ph=True),
                 _row("ECOM-2026-09-EBAY", "ebay", "Customer", "CUST-EBAY", 25.00, "stray receivable", ph=True)]
    run.allocations = pd.concat([a, pd.DataFrame(ebay)], ignore_index=True)
    run.revenue = pd.DataFrame([
        {"channel": "ebay", "customer_no": "CUST-EBAY", "account_no": "4xxxx-EBAY", "store_id": "Store01",
         "description": "eBay sales Store01", "amount": 400.00, "placeholder_account": True},
        {"channel": "ebay", "customer_no": "CUST-EBAY", "account_no": "4xxxx-EBAY", "store_id": "Store01",
         "description": "eBay sales Store01 (books)", "amount": 200.00, "placeholder_account": True},
        {"channel": "ebay", "customer_no": "CUST-EBAY", "account_no": "4xxxx-EBAY", "store_id": "Store02",
         "description": "eBay sales Store02", "amount": 400.00, "placeholder_account": True},
        {"channel": "ebay", "customer_no": "CUST-EBAY", "account_no": "4xxxx-SHIPINC", "store_id": None,
         "description": "eBay shipping", "amount": 60.00, "placeholder_account": True},
    ])
    return run


def test_v14_invoice_from_revenue_table(tmp_path):
    run = fake_run_v14()
    j = build_journal(run)
    invs, issues = invoice.build_invoices(j, run, {"ebay": {"label": "eBay", "doc_suffix": "EBAY"}})
    assert not issues and len(invs) == 1
    inv = invs[0]
    assert inv.customer_number == "CUST-EBAY" and inv.channel == "ebay" and inv.placeholder
    assert inv.total == Decimal("1060.00")
    by = {(l.account_no, l.store_code): l.unit_price for l in inv.lines}
    assert by == {("4xxxx-EBAY", "Store01"): Decimal("600.00"), ("4xxxx-EBAY", "Store02"): Decimal("400.00"),
                  ("4xxxx-SHIPINC", ""): Decimal("60.00")}
    p = invoice.invoices_json(j, invs)["salesInvoices"][0]
    assert p["customerNumber"] == "CUST-EBAY" and len(p["salesInvoiceLines"]) == 3
    assert all(l["lineType"] == "Account" and l["quantity"] == 1 for l in p["salesInvoiceLines"])
    assert {"code": "STORE", "valueCode": "Store02"} in p["salesInvoiceLines"][1]["dimensionSetLines"]
    # credit-signed revenue tables are accepted too
    run.revenue["amount"] = -run.revenue["amount"]
    assert invoice.build_invoices(j, run)[0][0].total == Decimal("1060.00")


def test_v14_channel_invariant():
    run = fake_run_v14()
    j = build_journal(run)
    invs, _ = invoice.build_invoices(j, run)
    rows = {b.customer_no: b for b in reconcile.channel_balances(invs, j)}
    b = rows["CUST-EBAY"]
    assert (b.invoice_total, b.fees, b.refunds, b.payouts) == (Decimal("1060.00"), Decimal("100.00"),
                                                               Decimal("50.00"), Decimal("700.00"))
    assert b.open_balance == Decimal("210.00") and b.difference == 0 and b.status == "ok"
    assert b.documents == ["ECOM-2026-09-EBAY"]
    # revenue left in the channel doc breaks the invariant and becomes a blocking error
    run2 = fake_run_v14(revenue_left_in_journal=True)
    j2 = build_journal(run2)
    b2 = {r.customer_no: r for r in reconcile.channel_balances(invoice.build_invoices(j2, run2)[0], j2)}["CUST-EBAY"]
    assert b2.status == "break" and b2.difference == Decimal("-25.00") and "Close v1.4" in b2.note
    assert reconcile.balance_exceptions([b2])[0]["severity"] == "error"


def test_v14_revenue_control_uses_invoice():
    run = fake_run_v14()
    j = build_journal(run)
    rows = {r.source: r for r in reconcile.reconcile(run.source_totals, j, None, {"4xxxx-EBAY"},
                                                     {"ECOM-2026-09-EBAY": Decimal("1000.00")})}
    assert rows["harmonized.fct_orders:ebay"].journal_total == Decimal("1000.00")
    assert rows["harmonized.fct_orders:ebay"].status == "ok"


def test_v14_api_reports_customer_balances(client):
    client.state["run"] = fake_run_v14()
    body = client.get("/api/close/run?month=2026-09").json()
    bal = body["customer_balances"]
    assert len(bal) == 1 and bal[0]["customer_no"] == "CUST-EBAY" and bal[0]["status"] == "ok"
    assert bal[0]["open_balance"] == 210.0
    assert body["invoices"][0]["total"] == 1060.0
    assert client.post("/api/close/approve?month=2026-09", json={}).status_code == 200
    assert client.get("/api/close/export?month=2026-09&kind=invoice_json").json()["salesInvoices"][0][
        "customerNumber"] == "CUST-EBAY"
