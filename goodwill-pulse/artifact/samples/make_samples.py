"""Build the sample upload files from the generated demo/history CSVs.  .venv/bin/python artifact/samples/make_samples.py"""
import glob, json, datetime as dt
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]; OUT = Path(__file__).resolve().parent
DAYS = ["2026-09-08", "2026-09-09", "2026-09-10"]         # Eastern business days (all match the warehouse for ShopGoodwill/GoodwillFinds)

up = pd.concat([pd.read_csv(f) for f in sorted(glob.glob(str(ROOT / "data/samples/history/paid_orders_*.csv")))])
up = up[up["Upright Order ID"].notna()].copy()
up["_t"] = pd.to_datetime(up["Paid At"])
up["_et"] = up._t.dt.tz_localize("America/Los_Angeles", ambiguous=True, nonexistent="shift_forward").dt.tz_convert("America/New_York").dt.strftime("%Y-%m-%d")
up = up[up._et.isin(DAYS)].sort_values("_t")
cm = pd.concat([pd.read_csv(f) for f in sorted(glob.glob(str(ROOT / "data/samples/history/orders2023*.csv")))]).drop_duplicates()
cm["_t"] = pd.to_datetime(cm["Order Date"])
cm["_et"] = cm._t.dt.tz_localize("UTC").dt.tz_convert("America/New_York").dt.strftime("%Y-%m-%d")
cm = cm[cm._et.isin(DAYS)].sort_values("_t")
UCOLS = list(pd.read_csv(sorted(glob.glob(str(ROOT / "data/samples/history/paid_orders_*.csv")))[0], nrows=0).columns)
CCOLS = list(cm.columns[:-2])

def clean(v):
    return None if pd.isna(v) else (v.item() if hasattr(v, "item") else v)

def upright_book(df, path, sheet="Paid orders"):
    wb = Workbook(); ws = wb.active; ws.title = sheet; ws.append(UCOLS)
    for c in ws[1]: c.font = Font(bold=True)
    for _, r in df.iterrows():
        row = [clean(r[c]) for c in UCOLS]; row[UCOLS.index("Paid At")] = r["_t"].to_pydatetime()
        ws.append(row)
    n = ws.max_row
    ws.append([None] * 9 + [round(float(df.Subtotal.sum()), 2), round(float(df["Shipping Total"].sum()), 2)] + [None] * 8)      # the SUM row Upright adds at the bottom (values, as in the export)
    for r in ws.iter_rows(min_row=2, max_row=n, min_col=19, max_col=19):
        for c in r: c.number_format = "mm/dd/yyyy hh:mm:ss"
    wb.save(path)

def cm_book(df, path):
    wb = Workbook(); ws = wb.active; ws.title = "Orders"; ws.append(CCOLS)
    for c in ws[1]: c.font = Font(bold=True)
    for _, r in df.iterrows():
        row = [clean(r[c]) for c in CCOLS]; row[0] = r["_t"].to_pydatetime(); ws.append(row)
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=1):
        for c in r: c.number_format = "yyyy-mm-dd hh:mm:ss"
    wb.save(path)

# (a) matches the warehouse
upright_book(up, OUT / "upright_paid_orders_match.xlsx")
cm_book(cm, OUT / "cashmonkey_orders_match.xlsx")
up.drop(columns=["_t", "_et"]).to_csv(OUT / "upright_paid_orders_match.csv", index=False)

# (b) a few amounts differ and one order is duplicated
b = up.copy().reset_index(drop=True)
sg = b[b.Channel == "Shopgoodwill"].index
b.loc[sg[3], "Subtotal"] = round(b.loc[sg[3], "Subtotal"] + 25.00, 2)
b.loc[sg[10], "Subtotal"] = round(b.loc[sg[10], "Subtotal"] - 7.50, 2)
gf = b[b.Channel == "GoodwillFinds"].index
b.loc[gf[1], "Final Value Fee"] = round(b.loc[gf[1], "Final Value Fee"] + 4.00, 2)
b = pd.concat([b, b.loc[[sg[5]]]]).reset_index(drop=True)                      # exact duplicate of one order
upright_book(b, OUT / "upright_paid_orders_differs.xlsx")

# (c) messy: renamed headers, title rows above, text dates, a blank row, a totals row, a test order, a bad amount, a bad date
rows = up[up.Channel.isin(["Shopgoodwill", "GoodwillFinds"]) & (up._et == DAYS[2])].head(14)
wb = Workbook(); ws = wb.active; ws.title = "Export"
ws.append(["Weekly sales export - Goodwill Michiana"]); ws.append([])
ws.append(["Order #", "Marketplace", "Date Ordered", "Item Subtotal", "Shipping Charged", "Marketplace Fee", "Buyer"])
for _, r in rows.iterrows():
    ws.append([str(r["Channel Order ID"]), r["Channel"], r["_t"].strftime("%m/%d/%Y %H:%M"), f"${r['Subtotal']:,.2f}", r["Shipping Total"], r["Final Value Fee"], r["Channel Buyer"]])
ws.append([None] * 7)
ws.append(["TEST-1", "Shopgoodwill", "09/10/2026 10:00", 99.0, 5.0, 1.0, "TEST"])
ws.append(["BAD-1", "Shopgoodwill", "not a date", 10.0, 1.0, 1.0, "x"])
ws.append(["BAD-2", "Shopgoodwill", "09/10/2026 11:00", "ten dollars", 1.0, 1.0, "y"])
ws.append(["Total", None, None, float(rows.Subtotal.sum()), float(rows["Shipping Total"].sum()), float(rows["Final Value Fee"].sum()), None])
wb.save(OUT / "messy_sales_export.xlsx")
print("written", sorted(p.name for p in OUT.glob("*.xlsx")), len(up), len(cm))
