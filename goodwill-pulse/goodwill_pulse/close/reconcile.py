"""Reconciliation: source file → loaded → rules output → journal, per source; plus the approval gate.

Control principle (PLAN 3.2, slide 40): every document balances, every line traces to a source, and any missing
report, failed rule or posting error is an exception with an owner that blocks approval until it is resolved or
overridden with a note.
"""
from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .bc_excel import FIXED_TS, MONEY_FMT
from .journal import Journal, money

TOL = Decimal("0.005")
BLOCKING = {"error"}


def _num(v) -> Decimal | None:
    if v is None:
        return None
    try:
        if v != v:                       # NaN
            return None
    except TypeError:
        pass
    return money(v)


@dataclass
class ReconRow:
    source: str
    rows: int | None
    file_total: Decimal | None
    loaded_total: Decimal | None
    rules_total: Decimal | None
    journal_total: Decimal | None
    journal_lines: int
    diff_file_loaded: Decimal | None
    diff_loaded_rules: Decimal | None
    diff_rules_journal: Decimal | None
    status: str                          # ok | explained | break | missing
    note: str = ""
    doc_no: str = ""
    measure: str = ""

    def to_dict(self) -> dict:
        f = lambda v: None if v is None else float(v)
        return {"source": self.source, "doc_no": self.doc_no, "measure": self.measure, "rows": self.rows,
                "file_total": f(self.file_total), "loaded_total": f(self.loaded_total),
                "rules_total": f(self.rules_total), "journal_total": f(self.journal_total),
                "journal_lines": self.journal_lines, "diff_file_loaded": f(self.diff_file_loaded),
                "diff_loaded_rules": f(self.diff_loaded_rules), "diff_rules_journal": f(self.diff_rules_journal),
                "status": self.status, "note": self.note}


def _diff(a, b):
    return None if a is None or b is None else a - b


def journal_measure(doc, measure: str, revenue_accounts: set[str] | None = None,
                    invoice_revenue: Decimal | None = None) -> Decimal:
    """Recompute the source's control measure from the journal document's own lines, independently of the rules
    engine, so rules output -> journal is a real check:
      postage paid ......... total debits             FedEx net ........... credit to the Vendor line
      item revenue ......... credits to the channel revenue accounts (net of debits to them)
      statement net ........ debit to the Customer line   jewelry reclass ..... total credits
      anything else ........ total debits
    """
    m = (measure or "").lower()
    lines = doc.lines
    if "fedex" in m:
        return -sum((l.amount for l in lines if l.account_type == "Vendor"), Decimal("0.00"))
    if "revenue" in m and "statement net" not in m and invoice_revenue is not None:
        return invoice_revenue                     # Close v1.4: revenue is booked by the AR invoice
    if "revenue" in m and "statement net" not in m:
        rev = revenue_accounts or {l.account_no for l in lines
                                   if l.account_type == "G/L Account" and l.store_code and l.credit}
        return -sum((l.amount for l in lines if l.account_no in rev), Decimal("0.00"))
    if "statement net" in m:
        return sum((l.amount for l in lines if l.account_type == "Customer"), Decimal("0.00"))
    if "jewelry" in m:
        return doc.total_credit
    return doc.total_debit


def invoice_measure(measure: str, doc_no: str, invoices: list) -> Decimal | None:
    """Close v1.4 measures that live on the AR invoices instead of the journal:
      'AR invoice: ...'                 -> total of that channel's invoice (doc suffix = channel doc)
      '... jewelry ... AR invoice(s)'   -> invoice lines on jewelry accounts, all invoices"""
    m = (measure or "").lower()
    if "ar invoice" not in m:
        return None
    if "jewelry" in m and not m.startswith("ar invoice"):
        return sum((l.amount for i in invoices for l in i.lines if "JEWEL" in l.account_no.upper()), Decimal("0.00"))
    suffix = (doc_no or "").rsplit("-", 1)[-1].upper()
    inv = next((i for i in invoices if i.external_document_number.upper().endswith("-" + suffix)), None)
    return inv.total if inv else Decimal("0.00")


def reconcile(source_totals: list[dict], journal: Journal, exceptions: list[dict] | None = None,
              revenue_accounts: set[str] | None = None,
              invoice_revenue: dict[str, Decimal] | None = None, invoices: list | None = None) -> list[ReconRow]:
    """`invoice_revenue`: doc_no -> AR invoice revenue for legacy 'item revenue' measures.
    `invoices`: the AR invoices, for Close v1.4 'AR invoice: ...' measures (see invoice_measure)."""
    invoice_revenue = invoice_revenue or {}
    """One row per CloseRun.source_totals entry. Journal side is found by the entry's `doc_no` (fallback: lines
    whose `source` equals the entry's source)."""
    exc_by_source: dict[str, list[dict]] = {}
    for e in exceptions or []:
        exc_by_source.setdefault(str(e.get("source") or ""), []).append(e)
    out: list[ReconRow] = []
    used_docs = set()
    for st in source_totals:
        src = str(st["source"])
        doc_no = str(st.get("doc_no") or "")
        measure = str(st.get("measure") or "")
        doc = journal.document(doc_no) if doc_no else None
        if doc is None and not doc_no:
            docs = [d for d in journal.documents if d.source == src]
            doc = docs[0] if len(docs) == 1 else None
        rows = st.get("rows")
        ft, lt, rt = _num(st.get("file_total")), _num(st.get("loaded_total")), _num(st.get("rules_total"))
        inv_m = invoice_measure(measure, doc_no, invoices) if invoices is not None else None
        if inv_m is not None:
            if doc is not None:
                used_docs.add(doc.document_no)
            jt, jl = inv_m, (len(doc.lines) if doc is not None else 0)
        elif doc is not None:
            used_docs.add(doc.document_no)
            jt, jl = journal_measure(doc, measure, revenue_accounts, invoice_revenue.get(doc.document_no)), len(doc.lines)
        elif doc_no in invoice_revenue and "revenue" in measure.lower():
            jt, jl = invoice_revenue[doc_no], 0
        else:
            jt, jl = (Decimal("0.00") if rt else None), 0
        d1, d2, d3 = _diff(ft, lt), _diff(lt, rt), _diff(rt, jt)
        diffs = [d for d in (d1, d2, d3) if d is not None and abs(d) > TOL]
        base = src.split(":")[0]
        related = exc_by_source.get(src, []) + (exc_by_source.get(base, []) if base != src else [])
        if not rows and ft is None and lt is None:
            missing = [e for e in related if "missing" in str(e.get("message", "")).lower()]
            status = "missing"
            note = missing[0]["message"] if missing else "No file received for this month."
        elif not diffs:
            status, note = "ok", ""
        else:
            explained = [e for e in related if e.get("amount") is not None]
            amounts = {money(e["amount"]).copy_abs() for e in explained}
            together = sum((money(e["amount"]).copy_abs() for e in explained), Decimal("0.00"))
            if explained and all(d.copy_abs() in amounts for d in diffs):
                status = "explained"
                note = "; ".join(e["message"] for e in explained if money(e["amount"]).copy_abs()
                                 in {d.copy_abs() for d in diffs})[:400]
            elif explained and all(abs(d.copy_abs() - together) <= TOL for d in diffs):
                status = "explained"
                note = (f"{diffs[0].copy_abs():,.2f} = {len(explained)} exceptions on this source "
                        f"(e.g. {explained[0]['message']})")[:400]
            else:
                status = "break"
                parts = []
                if d1 and abs(d1) > TOL:
                    parts.append(f"file vs loaded {d1:,.2f}")
                if d2 and abs(d2) > TOL:
                    parts.append(f"loaded vs rules {d2:,.2f}")
                if d3 and abs(d3) > TOL:
                    parts.append(f"rules vs journal {d3:,.2f}")
                note = "Unexplained difference: " + ", ".join(parts) + "."
                # an explaining exception may cover only part of the story; show it anyway
                if explained:
                    note += " Related: " + "; ".join(e["message"] for e in explained)[:300]
        out.append(ReconRow(src, int(rows) if rows is not None else None, ft, lt, rt, jt, jl, d1, d2, d3,
                            status, note, doc_no or (doc.document_no if doc else ""), measure))
    for d in journal.documents:
        if d.document_no not in used_docs:
            out.append(ReconRow(d.source, None, None, None, None, d.total_debit, len(d.lines), None, None, None,
                                "break", "Journal document with no source total; cannot trace it to a file.",
                                d.document_no, "total debits"))
    return out


def recon_exceptions(rows: list[ReconRow]) -> list[dict]:
    out = []
    for r in rows:
        if r.status == "break":
            out.append({"source": r.source, "rule_id": "reconciliation", "severity": "error", "message": r.note,
                        "amount": float(next((d for d in (r.diff_file_loaded, r.diff_loaded_rules,
                                                          r.diff_rules_journal) if d), 0)), "owner": "Accounting"})
    return out


# --- Close v1.4 invariant: per channel, invoice − fees − refunds − payouts = open customer balance -----------------
@dataclass
class ChannelBalance:
    channel: str
    customer_no: str
    documents: list[str]
    invoice_total: Decimal
    fees: Decimal
    refunds: Decimal
    payouts: Decimal
    other: Decimal                        # anything else posted against the customer (e.g. revenue left in a doc)
    customer_credits: Decimal             # net movement on the customer in the journal (negative = credited)
    open_balance: Decimal                 # invoice_total + customer_credits
    difference: Decimal                   # (invoice − fees − refunds − payouts) − open_balance
    status: str                           # ok | break
    note: str = ""

    def to_dict(self) -> dict:
        f = float
        return {"channel": self.channel, "customer_no": self.customer_no, "documents": self.documents,
                "invoice_total": f(self.invoice_total), "fees": f(self.fees), "refunds": f(self.refunds),
                "payouts": f(self.payouts), "other": f(self.other), "customer_credits": f(self.customer_credits),
                "open_balance": f(self.open_balance), "difference": f(self.difference), "status": self.status,
                "note": self.note}


def _is_refund(l) -> bool:
    return "REFUND" in l.account_no.upper() or "refund" in l.description.lower()


def channel_balances(invoices: list, journal: Journal) -> list[ChannelBalance]:
    """For each marketplace customer: the AR invoice (revenue) less what the channel journal documents settle
    against it. Fees, refunds and payouts are the debit lines in documents that credit the customer; any other
    line there (a revenue credit left in a channel doc, for instance) breaks the invariant and is reported."""
    by_cust: dict[str, object] = {i.customer_number: i for i in invoices}
    customers = {l.account_no for l in journal.lines if l.account_type == "Customer"} | set(by_cust)
    out = []
    for cust in sorted(customers):
        inv = by_cust.get(cust)
        docs = [d for d in journal.documents if any(l.account_type == "Customer" and l.account_no == cust
                                                    for l in d.lines)]
        fees = refunds = payouts = other = movement = Decimal("0.00")
        for d in docs:
            for l in d.lines:
                if l.account_type == "Customer":
                    if l.account_no == cust:
                        movement += l.amount
                    else:
                        other += l.amount
                elif l.account_type == "Bank Account" and l.debit:
                    payouts += l.debit
                elif l.account_type == "G/L Account" and l.debit:
                    if _is_refund(l):
                        refunds += l.debit
                    else:
                        fees += l.debit
                else:
                    other += l.amount
        inv_total = inv.total if inv else Decimal("0.00")
        open_bal = inv_total + movement
        diff = (inv_total - fees - refunds - payouts) - open_bal
        notes = []
        if abs(diff) > TOL:
            notes.append(f"Invariant off by {diff:,.2f}: the channel documents post {-other:,.2f} against the "
                         "customer that is not a fee, refund or payout"
                         + (" (revenue should be on the AR invoice only, Close v1.4)." if other < 0 else "."))
        if inv is None:
            notes.append("No AR invoice for this customer this month.")
        if any(not d.balanced for d in docs):
            notes.append("A channel document does not balance.")
        status = "break" if abs(diff) > TOL or any(not d.balanced for d in docs) else "ok"
        if status == "ok" and open_bal < -TOL:
            notes.append(f"Customer is overpaid by {-open_bal:,.2f} (payouts exceed this month's net revenue; "
                         "usually prior-month sales paid this month).")
        out.append(ChannelBalance(getattr(inv, "channel", "") or (docs[0].document_no.rsplit("-", 1)[-1].lower()
                                                                  if docs else ""),
                                  cust, [d.document_no for d in docs], inv_total, fees, refunds, payouts, other,
                                  movement, open_bal, diff, status, " ".join(notes)))
    return out


def balance_exceptions(rows: list[ChannelBalance]) -> list[dict]:
    return [{"source": "ar_invoice", "rule_id": "channel_invariant", "severity": "error",
             "message": f"{r.customer_no}: {r.note}", "amount": float(r.difference), "owner": "Accounting"}
            for r in rows if r.status == "break"]


# --- approval gate ------------------------------------------------------------------------------------------------
def exception_id(e: dict) -> str:
    key = "|".join(str(e.get(k, "")) for k in ("source", "rule_id", "severity", "message"))
    return hashlib.sha1(key.encode()).hexdigest()[:10]


@dataclass
class Approval:
    approved: bool
    blocked: bool
    blocking: list[dict] = field(default_factory=list)
    override: bool = False
    note: str = ""
    approver: str = ""
    approved_at: str = ""

    def to_dict(self) -> dict:
        return {"approved": self.approved, "blocked": self.blocked, "blocking": self.blocking,
                "override": self.override, "note": self.note, "approver": self.approver,
                "approved_at": self.approved_at}


def blocking_exceptions(exceptions: list[dict], resolved_ids: set[str] | None = None) -> list[dict]:
    resolved_ids = resolved_ids or set()
    return [e for e in exceptions if str(e.get("severity", "")).lower() in BLOCKING
            and str(e.get("status", "open")) == "open" and exception_id(e) not in resolved_ids]


def approve(exceptions: list[dict], *, override: bool = False, note: str = "", approver: str = "",
            now: str = "", resolved_ids: set[str] | None = None) -> Approval:
    """Approve the close. Blocked while error exceptions are open, unless overridden with a note."""
    blocking = blocking_exceptions(exceptions, resolved_ids)
    if blocking and not (override and note.strip()):
        return Approval(False, True, blocking)
    return Approval(True, bool(blocking), blocking, override=bool(blocking and override), note=note.strip(),
                    approver=approver, approved_at=now)


# --- one-page reconciliation report -------------------------------------------------------------------------------
RECON_COLUMNS = ["Source", "Document No.", "Measure", "Rows in file", "File total", "Loaded", "Rules output", "Journal", "Journal lines",
                 "Diff file→loaded", "Diff loaded→rules", "Diff rules→journal", "Status", "Note"]
STATUS_LABEL = {"ok": "OK", "explained": "Explained", "break": "BREAK", "missing": "Missing"}


def write_reconciliation_xlsx(month, rows: list[ReconRow], journal: Journal, exceptions: list[dict],
                              comparison: dict | None, approval: Approval | None, out_dir: Path,
                              balances: list[ChannelBalance] | None = None) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"Reconciliation {month:%Y-%m}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Reconciliation"
    ws.append([f"E-Commerce month-end close {month:%B %Y}: source to journal reconciliation (synthetic data)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Journal: {len(journal.documents)} documents, {len(journal.lines)} lines, debits "
               f"{journal.total_debit:,.2f}, credits {journal.total_credit:,.2f}, "
               f"{'balanced' if journal.balanced else 'NOT balanced'}. Rules version {journal.rules_version}."])
    if approval:
        ws.append([("Approved" if approval.approved else "Not approved")
                   + (f" by {approval.approver}" if approval.approver else "")
                   + (f" at {approval.approved_at}" if approval.approved_at else "")
                   + (f" with override: {approval.note}" if approval.override else "")])
    ws.append([])
    hdr = ws.max_row + 1
    ws.append(RECON_COLUMNS)
    for r in rows:
        d = r.to_dict()
        ws.append([r.source, r.doc_no, r.measure, r.rows, d["file_total"], d["loaded_total"], d["rules_total"], d["journal_total"],
                   r.journal_lines, d["diff_file_loaded"], d["diff_loaded_rules"], d["diff_rules_journal"],
                   STATUS_LABEL.get(r.status, r.status), r.note])
        if r.status in ("break", "missing"):
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="FAE3E0")
    for c in ws[hdr]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCEFEB")
    for rr in range(hdr + 1, ws.max_row + 1):
        for col in (5, 6, 7, 8, 10, 11, 12):
            ws.cell(rr, col).number_format = MONEY_FMT
    for i, name in enumerate(RECON_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = 50 if name == "Note" else 15
    ws.freeze_panes = ws.cell(hdr + 1, 1)

    if balances is not None:
        bs = wb.create_sheet("Customer balances")
        bs.append(["Customer", "Channel", "AR invoice", "Fees", "Refunds", "Payouts", "Other", "Open balance",
                   "Difference", "Status", "Note"])
        for b in balances:
            bs.append([b.customer_no, b.channel, float(b.invoice_total), float(b.fees), float(b.refunds),
                       float(b.payouts), float(b.other), float(b.open_balance), float(b.difference),
                       STATUS_LABEL.get(b.status, b.status), b.note])
        for rr in range(2, bs.max_row + 1):
            for col in range(3, 10):
                bs.cell(rr, col).number_format = MONEY_FMT
    ex = wb.create_sheet("Exceptions")
    ex.append(["ID", "Severity", "Source", "Rule", "Message", "Amount", "Owner", "Status"])
    for e in exceptions:
        ex.append([exception_id(e), e.get("severity"), e.get("source"), e.get("rule_id"), e.get("message"),
                   e.get("amount"), e.get("owner"), e.get("status", "open")])
    if comparison:
        wc = wb.create_sheet("Workbook comparison")
        wc.append(["Kind", "Document No.", "Account No.", "Department", "Store", "Description", "System",
                   "Workbook", "Difference", "Explanation"])
        for d in comparison.get("differences", []):
            wc.append([d["kind"], d["doc_no"], d.get("account_no"), d.get("department_code"), d.get("store_id"),
                       d.get("description"), d.get("system_amount"), d.get("workbook_amount"),
                       d.get("difference"), d.get("explanation")])
        wc.append([])
        wc.append([f"{comparison.get('matches', 0)} lines match; {len(comparison.get('differences', []))} differ."])
    for sheet in wb.worksheets[1:]:
        for c in sheet[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DCEFEB")
        sheet.freeze_panes = "A2"
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 18
    wb.properties.created = wb.properties.modified = FIXED_TS
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)
    return path
