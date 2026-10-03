"""Parallel-run check: compare our journal to the manual allocation workbook, line by line (PLAN 3.4 #4).

Workbook: `data/close_inputs/<yyyy-mm>/E-Commerce Allocation <yyyy-mm>.xlsx`, sheet `Journal Entries`, columns as the
CloseRun allocations minus rule_id/source_ref/placeholder_account (doc_no, posting_date, source, account_type,
account_no, department_code, store_id, description, amount). Header spelling is matched loosely
("Document No." / "doc_no" / "Doc No" all work).

Matching, per document:
  1. exact key (account_type, account_no, department_code, store_id, description) and same amount → match
  2. same key, different amount → `amount` difference (re-keyed number), with a digit-level explanation
  3. leftovers with the same amount but a different key → `coding` difference (wrong account/department/store)
  4. anything else → `missing_in_workbook` / `extra_in_workbook`
A workbook document that no longer balances is reported too.
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .journal import Journal, JournalLine, build_journal, money

SHEET = "Journal Entries"
ALIASES = {
    "doc_no": ["doc_no", "document_no", "document_no.", "doc_no.", "document_number", "document"],
    "posting_date": ["posting_date", "date"],
    "source": ["source"],
    "account_type": ["account_type"],
    "account_no": ["account_no", "account_no.", "account", "account_number", "gl_account"],
    "department_code": ["department_code", "department", "dept", "dept_code"],
    "store_id": ["store_id", "store", "store_code", "supplier"],
    "description": ["description"],
    "amount": ["amount"],
    "debit": ["debit", "debit_amount"],
    "credit": ["credit", "credit_amount"],
}
KEY_FIELDS = ("account_type", "account_no", "department_code", "store_id", "description")


def workbook_path(close_inputs_dir: Path, month: date) -> Path:
    return close_inputs_dir / f"{month:%Y-%m}" / f"E-Commerce Allocation {month:%Y-%m}.xlsx"


def _norm_header(h) -> str:
    return re.sub(r"[\s\-/]+", "_", str(h or "").strip().lower()).strip("_")


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_workbook(path: Path, sheet: str = SHEET) -> list[dict]:
    """Read the workbook's journal sheet into dicts with canonical keys (+ 'row' = Excel row number)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"{path.name} has no sheet '{sheet}' (sheets: {', '.join(wb.sheetnames)})")
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    header_idx, header = 0, None
    for i, r in enumerate(rows, 1):                     # first row that names a doc and an account column
        names = [_norm_header(c) for c in r]
        if any(n in ALIASES["doc_no"] for n in names) and any(n in ALIASES["account_no"] for n in names):
            header, header_idx = names, i
            break
    if header is None:
        raise ValueError(f"{path.name} / {sheet}: no header row with a document and account column")
    col = {}
    for canon, names in ALIASES.items():
        for j, h in enumerate(header):
            if h in names and canon not in col:
                col[canon] = j
    out = []
    for i, r in enumerate(rows, header_idx + 1):
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        get = lambda k: r[col[k]] if k in col and col[k] < len(r) else None
        doc = _s(get("doc_no"))
        if not doc or doc.upper().startswith("TOTAL"):
            continue
        if "amount" in col and get("amount") not in (None, ""):
            amt = money(get("amount"))
        else:
            amt = money(get("debit") or 0) - money(get("credit") or 0)
        pd_ = get("posting_date")
        out.append({"row": i, "doc_no": doc,
                    "posting_date": pd_.date() if isinstance(pd_, datetime) else pd_,
                    "source": _s(get("source")), "account_type": _s(get("account_type")),
                    "account_no": _s(get("account_no")), "department_code": _s(get("department_code")),
                    "store_id": _s(get("store_id")), "description": _s(get("description")), "amount": amt})
    wb.close()
    return out


def _sys_rec(l: JournalLine) -> dict:
    return {"doc_no": l.document_no, "account_type": l.account_type, "account_no": l.account_no,
            "department_code": l.department_code, "store_id": l.store_code, "description": l.description,
            "amount": l.amount, "line_no": l.line_no, "source_ref": l.source_ref}


def _key(r: dict) -> tuple:
    return tuple(r[k] for k in KEY_FIELDS)


def explain_amount(system: Decimal, workbook: Decimal) -> str:
    """Plain-English guess at how a number was mis-keyed."""
    a, b = f"{abs(system):.2f}", f"{abs(workbook):.2f}"
    diff = workbook - system
    if system == -workbook:
        return "Sign flipped: debit keyed as credit (or the reverse)."
    if len(a) == len(b):
        pos = [i for i in range(len(a)) if a[i] != b[i]]
        if len(pos) == 2 and pos[1] == pos[0] + 1 and a[pos[0]] == b[pos[1]] and a[pos[1]] == b[pos[0]]:
            return (f"Transposed digits: {b} keyed instead of {a} (difference {abs(diff):,.2f}, "
                    "divisible by 9, the signature of a transposition).")
        if len(pos) == 1:
            return f"One digit re-keyed wrong: {b} instead of {a}."
    if system and (abs(workbook) * 10 == abs(system) or abs(workbook) == abs(system) * 10):
        return f"Decimal point slipped: {b} instead of {a}."
    if abs(diff) % 9 == 0 and diff:
        return f"Workbook differs by {abs(diff):,.2f} (divisible by 9; a transposition is likely)."
    return f"Workbook differs by {diff:+,.2f}."


def compare(run_or_journal, workbook) -> dict:
    """Compare a CloseRun (or an already built Journal) with the manual workbook.

    `workbook` is a path to the .xlsx (sheet 'Journal Entries') or rows from `read_workbook`.
    Returns {matches, differences: [...], documents, system_lines, workbook_lines, status}.
    """
    journal = run_or_journal if isinstance(run_or_journal, Journal) else build_journal(run_or_journal)
    if isinstance(workbook, (str, Path)):
        path = Path(workbook)
        if not path.exists():
            return {"status": "no_workbook", "matches": 0, "differences": [], "documents": [],
                    "system_lines": len(journal.lines), "workbook_lines": 0,
                    "message": f"No manual workbook at {path.name}; parallel-run check skipped."}
        res = compare_lines(journal, read_workbook(path))
        res["workbook"] = path.name
        return res
    return compare_lines(journal, list(workbook))


def compare_lines(journal: Journal, workbook_rows: list[dict]) -> dict:
    sys_by_doc: dict[str, list[dict]] = defaultdict(list)
    for l in journal.lines:
        sys_by_doc[l.document_no].append(_sys_rec(l))
    wb_by_doc: dict[str, list[dict]] = defaultdict(list)
    for r in workbook_rows:
        wb_by_doc[r["doc_no"]].append(r)

    matches, differences, doc_summary = 0, [], []
    for doc in sorted(set(sys_by_doc) | set(wb_by_doc)):
        s_rows, w_rows = list(sys_by_doc.get(doc, [])), list(wb_by_doc.get(doc, []))
        doc_matches, doc_diffs = 0, []
        # 1. exact
        for s in list(s_rows):
            w = next((w for w in w_rows if _key(w) == _key(s) and w["amount"] == s["amount"]), None)
            if w:
                s_rows.remove(s); w_rows.remove(w); doc_matches += 1
        # 2. same key, amount differs
        for s in list(s_rows):
            w = next((w for w in w_rows if _key(w) == _key(s)), None)
            if w:
                s_rows.remove(s); w_rows.remove(w)
                doc_diffs.append(_diff("amount", doc, s, w, explain_amount(s["amount"], w["amount"])))
        # 3. same amount, coding differs
        for s in list(s_rows):
            w = next((w for w in w_rows if w["amount"] == s["amount"]), None)
            if w:
                s_rows.remove(s); w_rows.remove(w)
                changed = [f"{k.replace('_', ' ')} {s[k] or '(blank)'} → {w[k] or '(blank)'}"
                           for k in KEY_FIELDS if s[k] != w[k]]
                doc_diffs.append(_diff("coding", doc, s, w, "Same amount, coded differently: " + "; ".join(changed)
                                       + "."))
        for s in s_rows:
            doc_diffs.append(_diff("missing_in_workbook", doc, s, None,
                                   "Line is in the system journal but not in the workbook."))
        for w in w_rows:
            doc_diffs.append(_diff("extra_in_workbook", doc, None, w,
                                   "Line is in the workbook but the system has no matching line."))
        doc_diffs = _fold_follow_on(doc_diffs)
        wb_total = sum((w["amount"] for w in wb_by_doc.get(doc, [])), Decimal("0.00"))
        if doc in wb_by_doc and wb_total != 0:
            for d in doc_diffs:
                d["explanation"] += f" The workbook's {doc} is out of balance by {wb_total:,.2f}."
        matches += doc_matches
        differences += doc_diffs
        doc_summary.append({"doc_no": doc, "system_lines": len(sys_by_doc.get(doc, [])),
                            "workbook_lines": len(wb_by_doc.get(doc, [])), "matches": doc_matches,
                            "differences": len(doc_diffs), "workbook_balance": float(wb_total)})
    return {"matches": matches, "differences": differences, "documents": doc_summary,
            "system_lines": len(journal.lines), "workbook_lines": len(workbook_rows),
            "status": "match" if not differences else "differences"}


BALANCING_TYPES = ("Vendor", "Customer", "Bank Account")


def _fold_follow_on(diffs: list[dict]) -> list[dict]:
    """A re-keyed input flows through the workbook's formulas into the document's balancing line (vendor / customer /
    bank), so one keying error shows as two changed lines. Report it once, at the root cause, and attach the
    balancing line as `follow_on`."""
    inputs = [d for d in diffs if d["kind"] == "amount" and d["account_type"] not in BALANCING_TYPES]
    out = []
    for d in diffs:
        root = None
        if d["kind"] == "amount" and d["account_type"] in BALANCING_TYPES:
            root = next((i for i in inputs if abs(i["difference"] + d["difference"]) < 0.005
                         and "follow_on" not in i), None)
        if root is None:
            out.append(d)
            continue
        root["follow_on"] = {k: d[k] for k in ("account_type", "account_no", "description", "system_amount",
                                                "workbook_amount", "difference", "workbook_row")}
        root["explanation"] += (f" The workbook's balancing line ({d['account_type']} {d['account_no']}, row "
                                f"{d['workbook_row']}) follows the wrong input: {d['workbook_amount']:,.2f} instead of "
                                f"{d['system_amount']:,.2f}, so the document still balances and the error is "
                                "invisible in the workbook.")
    return out


def _diff(kind: str, doc: str, s: dict | None, w: dict | None, explanation: str) -> dict:
    base = s or w
    sa, wa = (s["amount"] if s else None), (w["amount"] if w else None)
    return {"kind": kind, "doc_no": doc, "account_type": base["account_type"], "account_no": base["account_no"],
            "department_code": base["department_code"], "store_id": base["store_id"],
            "description": base["description"],
            "workbook_account_no": w["account_no"] if w else None,
            "system_amount": float(sa) if sa is not None else None,
            "workbook_amount": float(wa) if wa is not None else None,
            "difference": float((wa or 0) - (sa or 0)),
            "workbook_row": w["row"] if w else None, "system_line_no": s["line_no"] if s else None,
            "source_ref": s["source_ref"] if s else None, "explanation": explanation}


def compare_with_file(journal: Journal, path: Path) -> dict:
    return compare(journal, path)


def comparison_exceptions(res: dict) -> list[dict]:
    """Workbook differences are warnings: the system is the source of truth, the workbook is the parallel run."""
    return [{"source": "workbook", "rule_id": "workbook_compare", "severity": "warning",
             "message": f"{d['doc_no']} {d['account_no']}: {d['explanation']}", "amount": d["difference"],
             "owner": "Accounting"} for d in res.get("differences", [])]
