"""Close stage 04 "Apply rules" (slide 40): config/close_rules.yaml -> balanced journal lines (CONTRACT section 5).

    run_close(month, harmonized_path, finance_path, rules_path) -> CloseRun

One journal document per source (ECOM-2026-09-SHIPPING, -FEDEX, -SGW, -EBAY, -AMAZON, -GWF, -GWB).
Close v1.4: revenue is booked ONCE, by the AR invoice -> ``CloseRun.revenue`` (gross sales incl. S&H by channel, store
and revenue account, jewelry already moved to the jewelry account). Channel documents carry no revenue lines: only
fees, refunds and payout settlement against the marketplace customer. Amount > 0 = debit, < 0 = credit; every document sums to 0 (all arithmetic in integer cents). Every line carries
``source_ref`` ('<table>:<key>'), so it traces back to a bank line, an invoice, a statement line or a harmonized slice.
Both databases are opened read_only. Problems become ``exceptions`` (with an owner), never silent drops.

CLI:  python -m goodwill_pulse.close.rules --month 2026-09 [--workbook]
"""
from __future__ import annotations

import argparse
import hashlib
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import duckdb
import pandas as pd
import yaml

from goodwill_pulse.close import enrich
from goodwill_pulse.config import CONFIG_DIR, DATA_DIR

DEFAULT_RULES = CONFIG_DIR / "close_rules.yaml"
DEFAULT_HARMONIZED = DATA_DIR / "harmonized.duckdb"
DEFAULT_FINANCE = DATA_DIR / "sources" / "finance.duckdb"

REVENUE_COLUMNS = ["channel", "customer_no", "store_id", "account_no", "amount", "placeholder_account", "description",
                   "source_ref"]
ALLOC_COLUMNS = ["doc_no", "posting_date", "rule_id", "source", "account_type", "account_no", "department_code",
                 "store_id", "description", "amount", "placeholder_account", "source_ref"]


@dataclass
class CloseRun:
    month: date
    rules_version: str
    allocations: pd.DataFrame
    exceptions: list[dict] = field(default_factory=list)
    source_totals: list[dict] = field(default_factory=list)
    # Close v1.4: AR invoice lines. One row per (channel, customer_no, store_id, account_no); amount > 0 = invoice
    # amount (credit to revenue when the invoice posts). Columns: REVENUE_COLUMNS.
    revenue: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=REVENUE_COLUMNS))

    def customer_balances(self) -> pd.DataFrame:
        """Per marketplace customer: invoice - fees - refunds - payouts = open balance (Close v1.4 invariant)."""
        rows = []
        a = self.allocations
        for (ch, cust), g in self.revenue.groupby(["channel", "customer_no"], sort=False):
            docs = set(a.loc[a.account_no == cust, "doc_no"])
            d = a[a.doc_no.isin(docs) & (a.amount > 0)]
            fees = round(d.loc[d.account_no.str.contains("FEES"), "amount"].sum(), 2)
            refunds = round(d.loc[d.account_no.str.contains("REFUNDS"), "amount"].sum(), 2)
            payouts = round(d.loc[d.account_type == "Bank Account", "amount"].sum(), 2)
            inv = round(g.amount.sum(), 2)
            rows.append({"channel": ch, "customer_no": cust, "invoice_total": inv, "fees": fees, "refunds": refunds,
                         "payouts": payouts, "open_balance": round(inv - fees - refunds - payouts, 2) + 0.0})
        return pd.DataFrame(rows, columns=["channel", "customer_no", "invoice_total", "fees", "refunds", "payouts",
                                           "open_balance"])

    def doc_totals(self) -> pd.DataFrame:
        """Per document: lines, debits, credits, balance (0.00 when balanced)."""
        a = self.allocations
        if a.empty:
            return pd.DataFrame(columns=["doc_no", "lines", "debits", "credits", "balance"])
        g = a.groupby("doc_no", sort=False)["amount"]
        return pd.DataFrame({
            "lines": g.size(),
            "debits": g.apply(lambda s: round(s[s > 0].sum(), 2)),
            "credits": g.apply(lambda s: round(-s[s < 0].sum(), 2)),
            "balance": g.apply(lambda s: round(s.sum(), 2) + 0.0),
        }).reset_index()

    @property
    def blocking(self) -> list[dict]:
        return [e for e in self.exceptions if e["severity"] == "error"]


# ----------------------------------------------------------------------------------------------------------- helpers
def cents(x) -> int:
    return 0 if x is None or (isinstance(x, float) and pd.isna(x)) else int(round(float(x) * 100))


def dollars(c: int) -> float:
    return round(c / 100, 2)


def allocate(total_cents: int, weights: dict) -> dict:
    """Split ``total_cents`` over keys pro rata to ``weights`` (largest remainder; sums exactly)."""
    keys = list(weights)
    wsum = sum(weights.values())
    if not keys:
        return {}
    if wsum == 0:
        out = {k: 0 for k in keys}
        out[keys[0]] = total_cents
        return out
    raw = {k: total_cents * weights[k] / wsum for k in keys}
    out = {k: int(raw[k] // 1) for k in keys}
    rest = total_cents - sum(out.values())
    for k in sorted(keys, key=lambda k: (-(raw[k] - out[k]), str(k)))[:rest]:
        out[k] += 1
    return out


def rules_version(rules_path: Path) -> str:
    text = Path(rules_path).read_bytes()
    cfg = yaml.safe_load(text)
    return f"{cfg.get('version', '0')}+{hashlib.sha256(text).hexdigest()[:8]}"


def balance_exceptions(allocations: pd.DataFrame, owner: str = "Controller (close owner)") -> list[dict]:
    """One 'error' exception per document whose lines do not sum to zero."""
    out = []
    if allocations.empty:
        return out
    for doc, amt in allocations.groupby("doc_no", sort=False)["amount"]:
        bal = sum(cents(a) for a in amt)
        if bal != 0:
            out.append({"source": doc, "rule_id": "balance_check", "severity": "error",
                        "message": f"{doc} is not balanced: debits - credits = {dollars(bal):,.2f}",
                        "amount": dollars(bal), "owner": owner})
    return out


class _Ctx:
    def __init__(self, month: date, cfg: dict, hcon, fcon):
        self.month = month
        self.start, self.end = enrich.month_bounds(month)
        self.cfg = cfg
        self.hcon = hcon
        self.fcon = fcon
        self.lines: list[dict] = []
        self.revenue: list[dict] = []
        self.exceptions: list[dict] = []
        self.totals: list[dict] = []
        self.ym = enrich.period_label(month)
        self._lookup = None

    # -- config
    def owner(self, rule) -> str:
        return self.cfg.get("owners", {}).get(rule.get("owner"), rule.get("owner") or "")

    def doc_no(self, suffix: str) -> str:
        return self.cfg.get("doc_no_format", "ECOM-{yyyy}-{mm}-{suffix}").format(
            yyyy=f"{self.month.year:04d}", mm=f"{self.month.month:02d}", suffix=suffix)

    def account(self, key_or_spec) -> dict:
        if isinstance(key_or_spec, dict):
            return key_or_spec
        return self.cfg["accounts"][key_or_spec]

    def channel_acct(self, channel: str, which: str) -> dict:
        ch = self.cfg["channels"][channel]
        if which == "revenue":
            return {"account_type": "G/L Account", "account_no": ch["revenue"], "placeholder": True}
        return {"account_type": "Customer", "account_no": ch["customer"], "placeholder": True}

    @property
    def lookup(self) -> dict:
        if self._lookup is None:
            self._lookup = enrich.sku_store_lookup(self.hcon) if self.hcon is not None else {}
        return self._lookup

    # -- output
    def line(self, doc, rule, source, acct, amount_c: int, description, source_ref, store=None, dept=None):
        if amount_c == 0:
            return
        a = self.account(acct)
        self.lines.append({
            "doc_no": doc, "posting_date": self.end, "rule_id": rule["id"], "source": source,
            "account_type": a["account_type"], "account_no": str(a["account_no"]),
            "department_code": dept, "store_id": store, "description": description,
            "amount": dollars(amount_c), "placeholder_account": bool(a.get("placeholder", False)),
            "source_ref": source_ref,
        })

    def rev(self, channel, store, acct, c: int, description, source_ref):
        """One AR-invoice revenue slice. acct: 'revenue' (channel revenue account) or an accounts key."""
        if c == 0:
            return
        if acct == "revenue":
            a = {"account_no": self.cfg["channels"][channel]["revenue"], "placeholder": True}
        else:
            a = self.account(acct)
        self.revenue.append({"channel": channel, "customer_no": self.cfg["channels"][channel]["customer"],
                             "store_id": store, "account_no": str(a["account_no"]), "kind": acct, "c": int(c),
                             "placeholder_account": bool(a.get("placeholder", False)),
                             "description": description, "source_ref": source_ref})

    def exc(self, source, rule, severity, message, amount=None):
        self.exceptions.append({"source": source, "rule_id": rule["id"] if isinstance(rule, dict) else rule,
                                "severity": severity, "message": message,
                                "amount": None if amount is None else round(float(amount), 2),
                                "owner": self.owner(rule) if isinstance(rule, dict) else rule})

    def total(self, source, doc, measure, rows, file_c, loaded_c, rules_c):
        self.totals.append({"source": source, "doc_no": doc, "measure": measure, "rows": int(rows),
                            "file_total": None if file_c is None else dollars(file_c),
                            "loaded_total": None if loaded_c is None else dollars(loaded_c),
                            "rules_total": None if rules_c is None else dollars(rules_c)})

    def doc_sum(self, doc) -> int:
        return sum(cents(l["amount"]) for l in self.lines if l["doc_no"] == doc)

    def has_table(self, con, table) -> bool:
        if con is None:
            return False
        return bool(con.execute("SELECT count(*) FROM information_schema.tables WHERE table_name = ?",
                                [table]).fetchone()[0])

    def missing(self, source, rule, doc, measure, what):
        self.exc(source, rule, "error", f"Missing source: {what} for {self.ym}")
        self.total(source, doc, measure, 0, None, None, None)


# ------------------------------------------------------------------------------------------------- rule: shipping
def rule_shipping_bank(ctx: _Ctx, rule: dict):
    doc, source = ctx.doc_no(rule["doc_suffix"]), rule["source"]
    measure = "postage paid from bank 0101"
    if not ctx.has_table(ctx.fcon, "bank_0101"):
        return ctx.missing(source, rule, doc, measure, "1st Source bank activity (acct 0101)")
    df = ctx.fcon.execute(
        "SELECT line_id, post_date, description, CAST(round(amount*100) AS BIGINT) AS c FROM bank_0101 "
        "WHERE account_no = ? AND post_date BETWEEN ? AND ? ORDER BY post_date, line_id",
        [str(rule.get("bank_account", "0101")), ctx.start, ctx.end]).df()
    vendors = {k.upper(): v for k, v in rule["vendors"].items()}

    def vendor_of(desc):
        d = str(desc).upper()
        return next((v for k, v in vendors.items() if d.startswith(k)), None)

    df["vendor"] = df["description"].map(vendor_of)
    df = df[df["vendor"].notna()]
    if df.empty:
        return ctx.missing(source, rule, doc, measure, "postage lines (OSM / Pitney Bowes / EasyPost) in bank 0101")
    dept = rule.get("department")
    per_vendor = defaultdict(int)
    for r in df.itertuples():
        ctx.line(doc, rule, source, rule["debit"], -r.c, f"{r.vendor} postage {r.post_date:%m/%d} ({r.description})",
                 f"bank_0101:{r.line_id}", dept=dept)
        per_vendor[r.vendor] += r.c
    for vendor, c in per_vendor.items():
        ctx.line(doc, rule, source, rule["credit"], c, f"{vendor} postage paid from 1st Source 0101, {ctx.ym}",
                 f"bank_0101:{ctx.ym}/{vendor}")
    file_c = -int(df["c"].sum())
    ctx.total(source, doc, measure, len(df), file_c, file_c,
              sum(cents(l["amount"]) for l in ctx.lines if l["doc_no"] == doc and l["amount"] > 0))


# ---------------------------------------------------------------------------------------------------- rule: fedex
def rule_fedex_net(ctx: _Ctx, rule: dict):
    doc, source = ctx.doc_no(rule["doc_suffix"]), rule["source"]
    measure = "FedEx charges net of refunds"
    if not ctx.has_table(ctx.fcon, "fedex_invoices"):
        return ctx.missing(source, rule, doc, measure, "FedEx invoices")
    dept = str(rule["department"])
    fx = ctx.fcon.execute(
        "SELECT line_id, invoice_no, invoice_date, line_type, tracking, CAST(round(amount*100) AS BIGINT) AS c, "
        "orig_invoice_no, refund_ref FROM fedex_invoices ORDER BY invoice_date, line_id").df()
    in_month = fx[(fx.invoice_date.dt.date >= ctx.start) & (fx.invoice_date.dt.date <= ctx.end)] if not fx.empty else fx
    charges = in_month[in_month.line_type == "CHARGE"]
    if charges.empty:
        return ctx.missing(source, rule, doc, measure, "FedEx invoice charges")

    charge_c = 0
    for inv, g in charges.groupby("invoice_no", sort=False):
        c = int(g["c"].sum())
        charge_c += c
        ctx.line(doc, rule, source, rule["expense"], c,
                 f"FedEx invoice {inv} ({g.invoice_date.iloc[0]:%m/%d/%Y}, {len(g)} shipments)",
                 f"fedex_invoices:{inv}", dept=dept)

    # refunds arrive as BNKDEPOSIT lines in bank 0101 -> net them (by bank date in month)
    bank = pd.DataFrame(columns=["line_id", "post_date", "description", "c", "reference"])
    if ctx.has_table(ctx.fcon, "bank_0101"):
        bank = ctx.fcon.execute(
            "SELECT line_id, post_date, description, CAST(round(amount*100) AS BIGINT) AS c, reference FROM bank_0101 "
            "WHERE type = ? AND upper(description) LIKE ? ORDER BY post_date, line_id",
            [rule.get("refund_bank_type", "BNKDEPOSIT"), f"%{rule.get('refund_match', 'FEDEX').upper()}%"]).df()
    if not bank.empty:
        bank["d"] = pd.to_datetime(bank["post_date"]).dt.date
    else:
        bank["d"] = []
    refunds = bank[(bank.d >= ctx.start) & (bank.d <= ctx.end)]
    credit_by_ref = {r.refund_ref: r for r in fx[fx.line_type == "REFUND"].itertuples()}
    refund_c = 0
    prior = []
    for r in refunds.itertuples():
        refund_c += r.c
        cm = credit_by_ref.get(r.reference)
        orig = f", orig. invoice {cm.orig_invoice_no}" if cm is not None else ""
        ctx.line(doc, rule, source, rule["expense"], -r.c,
                 f"FedEx refund {r.reference} deposited {r.d:%m/%d}{orig}", f"bank_0101:{r.line_id}", dept=dept)
        if cm is None:
            ctx.exc(source, rule, "warning", f"Bank 0101 deposit {r.line_id} '{r.description}' ({r.d}) has no "
                    f"matching FedEx credit memo; netted anyway", dollars(r.c))
        else:
            if cm.invoice_date.date() < ctx.start:
                prior.append((r.reference, r.c))
    if prior:
        ctx.exc(source, rule, "info", f"{len(prior)} FedEx refunds credited before {ctx.ym} were deposited in "
                f"{ctx.ym} and are netted here by bank date ({', '.join(r for r, _ in prior[:5])}"
                f"{', ...' if len(prior) > 5 else ''})", dollars(sum(c for _, c in prior)))
    ctx.line(doc, rule, source, rule["vendor"], -(charge_c - refund_c),
             f"FedEx {ctx.ym}: charges {dollars(charge_c):,.2f} net of BNKDEPOSIT refunds {dollars(refund_c):,.2f}",
             f"fedex_invoices:{ctx.ym}/net")

    # cross-month: credit memos dated in the month whose BNKDEPOSIT lands after month end
    horizon = ctx.end + timedelta(days=int(rule.get("cross_month_lookahead_days", 10)))
    dep_by_ref = {r.reference: r for r in bank.itertuples()}
    for cm in fx[fx.line_type == "REFUND"].itertuples():
        if not (ctx.start <= cm.invoice_date.date() <= ctx.end):   # credit memo dated in the close month
            continue
        dep = dep_by_ref.get(cm.refund_ref)
        if dep is None:
            ctx.exc(source, rule, "warning", f"FedEx credit {cm.refund_ref} ({cm.invoice_date:%Y-%m-%d}, invoice "
                    f"{cm.orig_invoice_no}) has no BNKDEPOSIT in bank 0101 yet: not netted in {ctx.ym}", dollars(-cm.c))
        elif ctx.end < dep.d <= horizon:
            ctx.exc(source, rule, "warning",
                    f"Cross-month FedEx refund: {cm.refund_ref} ${dollars(-cm.c):,.2f} for {ctx.ym} invoice "
                    f"{cm.orig_invoice_no} (credit dated {cm.invoice_date:%Y-%m-%d}) was deposited {dep.d:%Y-%m-%d}, "
                    f"after month end: not netted in {ctx.ym} (nets by bank date next month). Accrue or leave?",
                    dollars(-cm.c))
    file_c = int(in_month["c"].sum())        # FedEx file view: charges + credit memos dated in the month
    ctx.total(source, doc, measure, len(charges) + len(refunds), file_c, file_c, charge_c - refund_c)


# ---------------------------------------------------------------------------------------- channel docs (Close v1.4)
# Revenue is booked ONCE, by the AR invoice (CloseRun.revenue). Channel journal documents carry only:
#   fees        Dr marketplace fees        / Cr marketplace customer
#   refunds     Dr refunds (contra, store) / Cr marketplace customer
#   settlement  Dr Bank 0101 (payout)      / Cr marketplace customer
_REFUNDS_BY_STORE = """
    SELECT coalesce(s.store_id, '~none') AS store, CAST(round(sum(r.amount)*100) AS BIGINT) AS c
    FROM fct_refunds r LEFT JOIN (
        SELECT order_key, store_id FROM fct_order_lines
        QUALIFY row_number() OVER (PARTITION BY order_key ORDER BY sale_price DESC, line_no) = 1) s
      USING (order_key)
    WHERE r.channel = ? AND r.business_date BETWEEN ? AND ? GROUP BY 1 ORDER BY 1"""


def _fees(ctx: _Ctx, rule, doc, src, channel, fee_c: int, desc: str, ref: str):
    ctx.line(doc, rule, src, "marketplace_fees", fee_c, desc, ref)
    ctx.line(doc, rule, src, ctx.channel_acct(channel, "customer"), -fee_c, f"{desc} withheld", ref)


def _refunds(ctx: _Ctx, rule, doc, src, channel, by_store: dict, desc: str, ref: str) -> int:
    tot = 0
    for store, c in sorted(by_store.items(), key=lambda kv: str(kv[0])):
        st = None if store == "~none" else store
        ctx.line(doc, rule, src, "refunds", int(c), f"{desc} - {st or 'NO STORE'}", f"{ref}/{st or 'null'}", store=st)
        tot += int(c)
    ctx.line(doc, rule, src, ctx.channel_acct(channel, "customer"), -tot, f"{desc} owed back to buyers", ref)
    return tot


def _settlement(ctx: _Ctx, rule, doc, channel) -> int:
    """Payouts that landed in bank 0101 during the month: Dr Bank 0101 per deposit / Cr customer."""
    ch = ctx.cfg["channels"][channel]
    pat = ch.get("payout_description")
    if not pat or not ctx.has_table(ctx.fcon, "bank_0101"):
        return 0
    dep = ctx.fcon.execute(
        "SELECT line_id, post_date, reference, CAST(round(amount*100) AS BIGINT) AS c FROM bank_0101 "
        "WHERE upper(description) LIKE ? AND post_date BETWEEN ? AND ? ORDER BY post_date, line_id",
        [f"{pat.upper()}%", ctx.start, ctx.end]).fetchall()
    tot = 0
    for line_id, d, ref, c in dep:
        ctx.line(doc, rule, "bank_0101", "bank_0101", int(c), f"{ch['label']} payout {ref} deposited {d:%m/%d}",
                 f"bank_0101:{line_id}")
        tot += int(c)
    ctx.line(doc, rule, "bank_0101", ctx.channel_acct(channel, "customer"), -tot,
             f"{ch['label']} payouts received in 1st Source 0101, {ctx.ym} ({len(dep)} deposits)",
             f"bank_0101:{ctx.ym}/{channel}/payouts")
    if not dep:
        ctx.exc("bank_0101", rule, "info", f"No {ch['label']} payout deposits in bank 0101 during {ctx.ym}")
    return tot


def _detail(ctx: _Ctx, rule, doc, channel, d0: date, d1: date, tag: str = "") -> dict:
    """Invoice revenue (by store) + fee / refund journal lines from harmonized detail, one channel and date range."""
    src = rule["source"]
    pre = f"{tag} " if tag else ""
    label = ctx.cfg["channels"][channel]["label"]
    span = f"{d0:%m/%d}-{d1:%m/%d}"
    base = f"harmonized.fct_orders:{channel}/{d0}..{d1}"
    n_orders, sub_c, ship_c, fee_c = ctx.hcon.execute(
        "SELECT count(*), CAST(round(coalesce(sum(subtotal),0)*100) AS BIGINT), "
        "CAST(round(coalesce(sum(coalesce(shipping_charged,0)+coalesce(handling,0)),0)*100) AS BIGINT), "
        "CAST(round(coalesce(sum(marketplace_fees),0)*100) AS BIGINT) "
        "FROM fct_orders WHERE channel = ? AND business_date BETWEEN ? AND ?", [channel, d0, d1]).fetchone()
    lines = ctx.hcon.execute(
        "SELECT l.sku, l.store_id, CAST(round(l.sale_price*100) AS BIGINT) AS c FROM fct_order_lines l "
        "JOIN fct_orders o USING (order_key) WHERE o.channel = ? AND o.business_date BETWEEN ? AND ?",
        [channel, d0, d1]).df()
    line_c = 0
    if not lines.empty:
        lines = enrich.fill_store(lines)
        for store, c in lines.groupby(lines["store_id"].fillna("~none"), sort=True)["c"].sum().items():
            st = None if store == "~none" else store
            ctx.rev(channel, st, "revenue", int(c), f"{pre}{label} sales {span}",
                    f"harmonized.fct_order_lines:{channel}/{d0}..{d1}/{st or 'null'}")
            line_c += int(c)
            if st is None:
                ctx.exc(src, rule, "warning", f"{label}: {dollars(int(c)):,.2f} of sales have no store (SKU "
                        "unparseable); invoiced without store credit", dollars(int(c)))
        filled = int(lines["store_filled"].sum())
        if filled:
            ctx.exc(src, rule, "info", f"{label}: {filled} order lines had no store; filled from the SKU prefix")
    if sub_c != line_c:
        ctx.rev(channel, None, "revenue", sub_c - line_c, f"{pre}{label} order subtotal not on item lines",
                f"{base}/unallocated")
        ctx.exc(src, rule, "warning", f"{label}: order subtotal differs from item lines by "
                f"{dollars(sub_c - line_c):,.2f}; invoiced without store", dollars(sub_c - line_c))
    ctx.rev(channel, None, "shipping_income", ship_c, f"{pre}{label} shipping & handling charged {span}",
            f"{base}/shipping")
    _fees(ctx, rule, doc, src, channel, fee_c, f"{pre}{label} marketplace & payment fees {span}", f"{base}/fees")
    rf = dict(ctx.hcon.execute(_REFUNDS_BY_STORE, [channel, d0, d1]).fetchall())
    ref_c = _refunds(ctx, rule, doc, src, channel, rf, f"{pre}{label} refunds {span}",
                     f"harmonized.fct_refunds:{channel}/{d0}..{d1}")
    return {"orders": n_orders, "subtotal": sub_c, "shipping": ship_c, "fees": fee_c, "refunds": ref_c}


def _channel_total(ctx: _Ctx, source, doc, channel, rows, file_c):
    inv = sum(r["c"] for r in ctx.revenue if r["channel"] == channel)
    ctx.total(source, doc, "AR invoice: gross sales incl. S&H", rows, file_c, file_c, inv)


def rule_channel_revenue(ctx: _Ctx, rule: dict):
    channel = rule["channel"]
    doc, source = ctx.doc_no(rule["doc_suffix"]), rule["source"]
    n = ctx.hcon.execute("SELECT count(*) FROM fct_orders WHERE channel = ? AND business_date BETWEEN ? AND ?",
                         [channel, ctx.start, ctx.end]).fetchone()[0]
    if n == 0:
        return ctx.missing(source, rule, doc, "AR invoice: gross sales incl. S&H",
                           f"{ctx.cfg['channels'][channel]['label']} orders")
    s = _detail(ctx, rule, doc, channel, ctx.start, ctx.end)
    _settlement(ctx, rule, doc, channel)
    _channel_total(ctx, f"{source}:{channel}", doc, channel, s["orders"], s["subtotal"] + s["shipping"])


# ---------------------------------------------------------------------------------------------- rule: ShopGoodwill
def rule_sgw_periods(ctx: _Ctx, rule: dict):
    channel = rule["channel"]
    doc, source = ctx.doc_no(rule["doc_suffix"]), rule["source"]
    label = ctx.cfg["channels"][channel]["label"]
    n = ctx.hcon.execute("SELECT count(*) FROM fct_orders WHERE channel = ? AND business_date BETWEEN ? AND ?",
                         [channel, ctx.start, ctx.end]).fetchone()[0]
    if n == 0:
        return ctx.missing(source, rule, doc, "AR invoice: gross sales incl. S&H", f"{label} orders")
    tol = cents(rule.get("statement_tolerance", 1.0))
    file_c = 0
    for p, spec in sorted((int(k), v) for k, v in rule["periods"].items()):
        lo, hi = spec["days"]
        d0 = date(ctx.month.year, ctx.month.month, lo)
        d1 = min(ctx.end, date(ctx.month.year, ctx.month.month, min(hi, ctx.end.day)))
        tag, span = f"P{p}", f"{d0:%m/%d}-{d1:%m/%d}"
        st = ctx.hcon.execute(
            "SELECT payout_key, paid_on, CAST(round(gross*100) AS BIGINT), CAST(round(fees*100) AS BIGINT), "
            "CAST(round(refunds*100) AS BIGINT), CAST(round(net*100) AS BIGINT) FROM fct_payouts "
            "WHERE channel = ? AND period_start = ? ORDER BY payout_key LIMIT 1", [channel, d0]).fetchone()
        det = ctx.hcon.execute(
            "SELECT CAST(round(coalesce(sum(subtotal),0)*100) AS BIGINT), "
            "CAST(round(coalesce(sum(coalesce(shipping_charged,0)+coalesce(handling,0)),0)*100) AS BIGINT) "
            "FROM fct_orders WHERE channel = ? AND business_date BETWEEN ? AND ?", [channel, d0, d1]).fetchone()
        mode = spec["mode"]
        if mode == "periodic_only" and st is None:
            ctx.exc(source, rule, "error", f"Missing source: {label} Period {p} periodic statement "
                    f"({d0}..{d1}); posted from detail reports instead")
            mode = "all_reports"
        if mode == "all_reports":
            s = _detail(ctx, rule, doc, channel, d0, d1, tag=tag)
            file_c += s["subtotal"] + s["shipping"]
            if st is not None and abs(st[2] - s["subtotal"]) > tol and abs(st[2] - s["subtotal"] - s["shipping"]) > tol:
                ctx.exc(source, rule, "info", f"{label} Period {p} statement gross {dollars(st[2]):,.2f} differs "
                        f"from detail {dollars(s['subtotal'] + s['shipping']):,.2f}",
                        dollars(st[2] - s["subtotal"] - s["shipping"]))
            continue
        key, paid_on, gross_c, fees_c, ref_c, net_c = st
        file_c += gross_c
        if gross_c - fees_c - ref_c != net_c:
            ctx.exc(source, rule, "warning", f"{label} Period {p} statement {key} does not foot: gross - fees - "
                    f"refunds = {dollars(gross_c - fees_c - ref_c):,.2f}, net = {dollars(net_c):,.2f}",
                    dollars(gross_c - fees_c - ref_c - net_c))
        if abs(gross_c - det[0]) > tol and abs(gross_c - det[0] - det[1]) > tol:
            ctx.exc(source, rule, "warning", f"{label} Period {p} statement gross {dollars(gross_c):,.2f} differs "
                    f"from detail sales {dollars(det[0] + det[1]):,.2f}", dollars(gross_c - det[0] - det[1]))
        # store credit: statement gross (incl. S&H) split pro rata to the period's detail sales by store + detail S&H
        w = ctx.hcon.execute(
            "SELECT l.sku, l.store_id, CAST(round(l.sale_price*100) AS BIGINT) AS c FROM fct_order_lines l "
            "JOIN fct_orders o USING (order_key) WHERE o.channel = ? AND o.business_date BETWEEN ? AND ?",
            [channel, d0, d1]).df()
        weights = {}
        if not w.empty:
            w = enrich.fill_store(w)
            weights = w.groupby(w["store_id"].fillna("~none"), sort=True)["c"].sum().astype(int).to_dict()
        if det[1]:
            weights["~ship"] = int(det[1])
        ref = f"harmonized.fct_payouts:{key}"
        for store, c in allocate(gross_c, weights or {"~none": 1}).items():
            if store == "~ship":
                ctx.rev(channel, None, "shipping_income", c, f"{tag} {label} statement {span} shipping & handling share",
                        f"{ref}/shipping")
            else:
                stid = None if store == "~none" else store
                ctx.rev(channel, stid, "revenue", c, f"{tag} {label} periodic statement {span} gross",
                        f"{ref}/{stid or 'null'}")
        _fees(ctx, rule, doc, source, channel, fees_c, f"{tag} {label} statement commission & payment fees {span}",
              f"{ref}/fees")
        rw = dict(ctx.hcon.execute(_REFUNDS_BY_STORE, [channel, d0, d1]).fetchall())
        _refunds(ctx, rule, doc, source, channel, allocate(ref_c, rw or {"~none": 1}),
                 f"{tag} {label} statement refunds {span}", f"{ref}/refunds")
    _settlement(ctx, rule, doc, channel)
    _channel_total(ctx, source, doc, channel, n, file_c)


# ---------------------------------------------------------------------------------------------- rule: Goodwillbooks
def rule_gwb_statement(ctx: _Ctx, rule: dict):
    channel = rule["channel"]
    doc, source = ctx.doc_no(rule["doc_suffix"]), rule["source"]
    measure = "AR invoice: statement gross (SALE + SHIPPING)"
    smonth = enrich.period_label(enrich.add_months(ctx.month, int(rule.get("statement_month_offset", 0))))
    have = ctx.has_table(ctx.fcon, "gwb_statement")
    st = ctx.fcon.execute(
        "SELECT line_no, line_type, order_no, sku, description, CAST(round(amount*100) AS BIGINT) AS c "
        "FROM gwb_statement WHERE statement_month = ? ORDER BY line_no", [smonth]).df() if have else pd.DataFrame()
    label = ctx.cfg["channels"][channel]["label"]
    if st.empty:
        ctx.exc(source, rule, "error", f"Missing source: Goodwill Books payment statement {smonth} not received; "
                f"{ctx.ym} accrued from harmonized order detail instead")
        accrual = dict(rule, id=f"{rule['id']}_accrual", source="harmonized.fct_orders")
        n = ctx.hcon.execute("SELECT count(*) FROM fct_orders WHERE channel = ? AND business_date BETWEEN ? AND ?",
                             [channel, ctx.start, ctx.end]).fetchone()[0]
        if n:
            s = _detail(ctx, accrual, doc, channel, ctx.start, ctx.end, tag="ACCRUAL")
            _settlement(ctx, accrual, doc, channel)
            inv = sum(r["c"] for r in ctx.revenue if r["channel"] == channel)
            ctx.total(source, doc, measure, 0, None, s["subtotal"] + s["shipping"], inv)
        else:
            ctx.total(source, doc, measure, 0, None, None, None)
        return
    header = None
    if ctx.has_table(ctx.fcon, "gwb_statement_header"):
        header = ctx.fcon.execute("SELECT CAST(round(gross*100) AS BIGINT), CAST(round(net*100) AS BIGINT) "
                                  "FROM gwb_statement_header WHERE statement_month = ?", [smonth]).fetchone()
    lines_c = int(st["c"].sum())
    if header is None:
        ctx.exc(source, rule, "warning", f"Goodwill Books statement {smonth} has no total line to control against")
    elif header[1] != lines_c:
        ctx.exc(source, rule, "error", f"Goodwill Books statement {smonth} lines sum to {dollars(lines_c):,.2f} but "
                f"the statement total is {dollars(header[1]):,.2f}", dollars(lines_c - header[1]))
    st["store_id"] = st["sku"].map(lambda s: ctx.lookup.get(s) or enrich.store_from_sku(s))
    st["store_key"] = st["store_id"].fillna("~none")
    known = {"SALE", "SHIPPING", "FEE", "REFUND"}
    for lt, g in st[~st.line_type.isin(known)].groupby("line_type"):
        ctx.exc(source, rule, "warning", f"Goodwill Books statement line type '{lt}' has no rule; treated as a fee",
                dollars(int(g["c"].sum())))
    ref = f"gwb_statement:{smonth}"
    for store, g in st[st.line_type == "SALE"].groupby("store_key", sort=True):
        stid = None if store == "~none" else store
        ctx.rev(channel, stid, "revenue", int(g["c"].sum()), f"{label} statement {smonth} sales",
                f"{ref}/SALE/{stid or 'null'} (lines {g.line_no.min()}-{g.line_no.max()})")
        if stid is None:
            ctx.exc(source, rule, "warning", f"{label}: {len(g)} statement sale lines have no store (SKU not found)",
                    dollars(int(g["c"].sum())))
    ship = st[st.line_type == "SHIPPING"]
    ctx.rev(channel, None, "shipping_income", int(ship["c"].sum()), f"{label} statement {smonth} shipping",
            f"{ref}/SHIPPING")
    fee_c = -int(st[~st.line_type.isin(["SALE", "SHIPPING", "REFUND"])]["c"].sum())
    _fees(ctx, rule, doc, source, channel, fee_c, f"{label} statement {smonth} platform fees", f"{ref}/FEE")
    rf = st[st.line_type == "REFUND"]
    _refunds(ctx, rule, doc, source, channel, (-rf.groupby("store_key")["c"].sum()).astype(int).to_dict(),
             f"{label} statement {smonth} refunds", f"{ref}/REFUND")
    _settlement(ctx, rule, doc, channel)
    inv = sum(r["c"] for r in ctx.revenue if r["channel"] == channel)
    gross_lines = int(st[st.line_type.isin(["SALE", "SHIPPING"])]["c"].sum())
    ctx.total(source, doc, measure, len(st), None if header is None else header[0], gross_lines, inv)


# ---------------------------------------------------------------------------------------------------- rule: jewelry
def rule_jewelry_reclass(ctx: _Ctx, rule: dict):
    """Jewelry sales are invoiced to the jewelry revenue account (per store) instead of the channel revenue account.
    Under Close v1.4 this changes CloseRun.revenue only; there is no JEWELRY journal document."""
    source = rule["source"]
    measure = "jewelry sales moved to jewelry revenue on the AR invoices"
    if not ctx.has_table(ctx.fcon, "jewelry_report"):
        return ctx.missing(source, rule, None, measure, "Jewelry Report")
    jr = ctx.fcon.execute(
        "SELECT line_no, channel, order_ref, sku, supplier, CAST(round(sale_price*100) AS BIGINT) AS c "
        "FROM jewelry_report WHERE report_month = ? ORDER BY line_no", [ctx.ym]).df()
    if jr.empty:
        return ctx.missing(source, rule, None, measure, "Jewelry Report rows")
    jr = enrich.label(enrich.fill_supplier(jr, ctx.lookup), source, ctx.month)
    known = jr.channel.isin(list(ctx.cfg["channels"]))
    for r in jr[jr.supplier_source == "unmatched"].itertuples():
        ctx.exc(source, rule, "error", f"Jewelry Report line {r.line_no}: SKU '{r.sku}' (order {r.order_ref}) not "
                f"found, Supplier cannot be filled; {dollars(r.c):,.2f} left in channel revenue, not jewelry", dollars(r.c))
    for r in jr[(jr.supplier_source != "unmatched") & ~known].itertuples():
        ctx.exc(source, rule, "error", f"Jewelry Report line {r.line_no}: unknown channel '{r.channel}'", dollars(r.c))
    ok = jr[(jr.supplier_source != "unmatched") & known]
    moved = 0
    for (ch, store), g in ok.groupby(["channel", "store_id"], sort=True):
        c = int(g["c"].sum())
        have = sum(r["c"] for r in ctx.revenue if r["channel"] == ch and r["store_id"] == store and r["kind"] == "revenue")
        if c > have:
            ctx.exc(source, rule, "warning", f"Jewelry {ch} {store}: {dollars(c):,.2f} on the report exceeds invoiced "
                    f"{ch} sales for the store ({dollars(have):,.2f}); not reclassified", dollars(c - have))
            continue
        ref = f"jewelry_report:{ctx.ym}/{ch}/{store}"
        ctx.rev(ch, store, "revenue", -c, f"Jewelry reclass ({len(g)} items)", ref)
        ctx.rev(ch, store, rule["credit"], c, f"Jewelry sales ({len(g)} items)", ref)
        moved += c
    ctx.total(source, None, measure, len(jr), int(jr["c"].sum()), int(jr["c"].sum()), moved)


HANDLERS = {
    "shipping_bank": rule_shipping_bank,
    "fedex_net": rule_fedex_net,
    "channel_revenue": rule_channel_revenue,
    "sgw_periods": rule_sgw_periods,
    "gwb_statement": rule_gwb_statement,
    "jewelry_reclass": rule_jewelry_reclass,
}


# --------------------------------------------------------------------------------------------------------- run_close
def _revenue_frame(ctx: _Ctx) -> pd.DataFrame:
    """Aggregate revenue slices to one AR invoice line per (channel, customer, store, account), in config order."""
    if not ctx.revenue:
        return pd.DataFrame(columns=REVENUE_COLUMNS)
    order = {ch: i for i, ch in enumerate(ctx.cfg["channels"])}
    agg: dict = {}
    for r in ctx.revenue:
        k = (r["channel"], r["customer_no"], r["store_id"], r["account_no"])
        e = agg.setdefault(k, {"c": 0, "refs": [], "ph": r["placeholder_account"]})
        e["c"] += r["c"]
        if r["source_ref"] not in e["refs"]:
            e["refs"].append(r["source_ref"])
    labels = {v["account_no"]: v.get("label", "") for v in ctx.cfg["accounts"].values()}
    rows = []
    for (ch, cust, store, acct), e in sorted(agg.items(), key=lambda kv: (order.get(kv[0][0], 99), kv[0][2] or "~",
                                                                           kv[0][3])):
        if e["c"] == 0:
            continue
        lab = ctx.cfg["channels"][ch]["label"]
        what = labels.get(acct, "").replace("PLACEHOLDER ", "") or "sales"
        rows.append({"channel": ch, "customer_no": cust, "store_id": store, "account_no": acct,
                     "amount": dollars(e["c"]), "placeholder_account": e["ph"],
                     "description": f"{lab} {ctx.ym} {what}" + (f" - {store}" if store else ""),
                     "source_ref": "; ".join(e["refs"])})
    df = pd.DataFrame(rows, columns=REVENUE_COLUMNS)
    df["store_id"] = pd.Series([v if isinstance(v, str) else None for v in df["store_id"]], index=df.index, dtype=object)
    return df


def _open(path) -> duckdb.DuckDBPyConnection | None:
    p = Path(path) if path else None
    if p is None or not p.exists():
        return None
    return duckdb.connect(str(p), read_only=True)


def run_close(month: date, harmonized_path=DEFAULT_HARMONIZED, finance_path=DEFAULT_FINANCE,
              rules_path=DEFAULT_RULES) -> CloseRun:
    """Apply every rule in ``rules_path`` for ``month`` (any date in the month). Never raises for data problems:
    they come back as exceptions. Raises FileNotFoundError only if the harmonized DB or rules file is missing."""
    month = date(month.year, month.month, 1)
    rules_path = Path(rules_path)
    cfg = yaml.safe_load(rules_path.read_text())
    hcon = _open(harmonized_path)
    if hcon is None:
        raise FileNotFoundError(f"harmonized DB not found: {harmonized_path}")
    fcon = _open(finance_path)
    try:
        ctx = _Ctx(month, cfg, hcon, fcon)
        if fcon is None:
            ctx.exc("finance", {"id": "finance_db", "owner": "close"}, "error",
                    f"Missing source: finance DB not found at {finance_path} (bank 0101, FedEx, jewelry, GWB)")
        for rule in cfg["rules"]:
            HANDLERS[rule["kind"]](ctx, rule)
        alloc = pd.DataFrame(ctx.lines, columns=ALLOC_COLUMNS)
        alloc["amount"] = alloc["amount"].astype(float)
        alloc["placeholder_account"] = alloc["placeholder_account"].astype(bool)
        for col in ("doc_no", "rule_id", "source", "account_type", "account_no", "department_code", "store_id",
                    "description", "source_ref"):
            alloc[col] = pd.Series([v if isinstance(v, str) else None for v in alloc[col]], index=alloc.index,
                                   dtype=object)
        revenue = _revenue_frame(ctx)
        neg = revenue[revenue.amount < 0]
        for r in neg.itertuples():
            ctx.exc("revenue", {"id": "ar_invoice", "owner": "revenue"}, "warning",
                    f"Negative invoice line {r.channel} {r.store_id} {r.account_no}: {r.amount:,.2f}", r.amount)
        exceptions = ctx.exceptions + balance_exceptions(alloc, cfg.get("owners", {}).get("close", "close"))
        n_ph = int(alloc["placeholder_account"].sum()) + int(revenue["placeholder_account"].sum())
        if n_ph:
            exceptions.append({"source": "close_rules.yaml", "rule_id": "placeholder_accounts", "severity": "info",
                               "message": f"{n_ph} journal / AR invoice lines post to PLACEHOLDER accounts (4xxxx/6xxxx/CUST-*); "
                                          "replace with Goodwill's chart of accounts before posting",
                               "amount": None, "owner": cfg.get("owners", {}).get("close", "close")})
        return CloseRun(month=month, rules_version=rules_version(rules_path), allocations=alloc,
                        exceptions=exceptions, source_totals=ctx.totals, revenue=revenue)
    finally:
        hcon.close()
        if fcon is not None:
            fcon.close()


# ------------------------------------------------------------------------------- synthetic MANUAL allocation workbook
WORKBOOK_COLUMNS = ["doc_no", "posting_date", "source", "account_type", "account_no", "department_code", "store_id",
                    "description", "amount"]
ORANGE = "FFC000"
BALANCING_TYPES = ("Vendor", "Customer", "Bank Account")


def _transpose(c: int, rng) -> int | None:
    """Swap two adjacent, different digits in the whole-dollar part of |c| cents (a classic re-keying error)."""
    sign, c = (-1 if c < 0 else 1), abs(c)
    d, cc = str(c // 100), c % 100
    pairs = [i for i in range(len(d) - 1) if d[i] != d[i + 1] and not (i == 0 and d[i + 1] == "0")]
    if not pairs:
        return None
    i = pairs[int(rng.integers(len(pairs)))]
    d2 = d[:i] + d[i + 1] + d[i] + d[i + 2:]
    return sign * (int(d2) * 100 + cc)


def plant_error(alloc: pd.DataFrame, seed: int = 7) -> tuple[pd.DataFrame, dict]:
    """Copy of the journal lines with ONE re-keyed input amount (transposed digits). The document's balancing line
    (vendor / customer / bank) follows the wrong input, as the workbook's formulas would, so the doc still balances."""
    import numpy as np
    rng = np.random.default_rng(seed)
    wb = alloc.reset_index(drop=True).copy()
    is_input = ~wb["account_type"].isin(BALANCING_TYPES) & (wb["amount"].abs() >= 100)
    pref = wb.index[is_input & wb["doc_no"].str.endswith("-FEDEX") & (wb["amount"] > 0)].tolist()
    cands = pref or wb.index[is_input].tolist()
    order = list(rng.permutation(cands)) if cands else []
    for idx in order:
        old = cents(wb.at[idx, "amount"])
        new = _transpose(old, rng)
        if new is None:
            continue
        doc = wb.at[idx, "doc_no"]
        bal = wb.index[(wb["doc_no"] == doc) & wb["account_type"].isin(BALANCING_TYPES)].tolist()
        if not bal:
            continue
        after = [j for j in bal if j > idx]
        b = after[0] if after else bal[-1]
        wb.at[idx, "amount"] = dollars(new)
        wb.at[b, "amount"] = dollars(cents(wb.at[b, "amount"]) - (new - old))
        return wb, {"row_index": int(idx), "balancing_index": int(b), "doc_no": doc,
                    "description": wb.at[idx, "description"], "correct": dollars(old), "keyed": dollars(new),
                    "delta": dollars(new - old), "balancing_description": wb.at[b, "description"],
                    "balancing_correct": dollars(cents(wb.at[b, "amount"]) + (new - old)),
                    "balancing_keyed": float(wb.at[b, "amount"])}
    raise ValueError("no journal line suitable for a planted re-keying error")


def write_manual_workbook(close_run: CloseRun, out_dir, seed: int = 7) -> Path:
    """Write the synthetic MANUAL allocation workbook (what the team keys by hand today, slide 39) for the close month,
    with exactly one planted re-keying error, plus ANSWER_KEY.md describing it. Returns the .xlsx path."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    ym = enrich.period_label(close_run.month)
    wbdf, err = plant_error(close_run.allocations, seed)
    wbdf = wbdf[WORKBOOK_COLUMNS]

    wb = Workbook()
    ws_in = wb.active
    ws_in.title = "Inputs"
    ws = wb.create_sheet("Journal Entries")
    bold = Font(bold=True)
    orange = PatternFill("solid", fgColor=ORANGE)

    ws_in.append([f"E-Commerce Allocation {ym} - INPUTS (orange cells are typed from the source reports)"])
    ws_in["A1"].font = bold
    ws_in.append(["Input #", "Document No.", "Source", "What was keyed", "Amount", "Feeds Journal Entries row"])
    for c in ws_in[2]:
        c.font = bold
    ws.append(WORKBOOK_COLUMNS)
    for c in ws[1]:
        c.font = bold
    err_cell = None
    n = 0
    for i, r in enumerate(wbdf.itertuples(index=False)):
        je_row = i + 2
        ws.append([r.doc_no, r.posting_date, r.source, r.account_type, r.account_no, r.department_code, r.store_id,
                   r.description, float(r.amount)])
        ws.cell(je_row, 2).number_format = "yyyy-mm-dd"
        ws.cell(je_row, 9).number_format = "#,##0.00"
        if r.account_type in BALANCING_TYPES:
            continue
        n += 1
        ws_in.append([n, r.doc_no, r.source, r.description, float(r.amount), f"'Journal Entries'!I{je_row}"])
        cell = ws_in.cell(ws_in.max_row, 5)
        cell.fill = orange
        cell.number_format = "#,##0.00"
        if i == err["row_index"]:
            err_cell = f"Inputs!E{ws_in.max_row}"
            err["je_cell"] = f"'Journal Entries'!I{je_row}"
    err["balancing_je_cell"] = f"'Journal Entries'!I{err['balancing_index'] + 2}"
    # step 06 today: the AR invoice is created from the Invoices tab (revenue by marketplace and store)
    ws_inv = wb.create_sheet("Invoices")
    ws_inv.append(["channel", "customer_no", "store_id", "account_no", "description", "amount"])
    for c in ws_inv[1]:
        c.font = bold
    for r in close_run.revenue.itertuples(index=False):
        ws_inv.append([r.channel, r.customer_no, r.store_id, r.account_no, r.description, float(r.amount)])
        ws_inv.cell(ws_inv.max_row, 6).number_format = "#,##0.00"
    for j, w in enumerate([14, 12, 10, 16, 60, 14], 1):
        ws_inv.column_dimensions[get_column_letter(j)].width = w
    err["input_cell"] = err_cell
    for sheet, widths in ((ws, [22, 12, 22, 14, 16, 10, 10, 70, 14]), (ws_in, [8, 24, 22, 70, 14, 28])):
        for j, w in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(j)].width = w
    from datetime import datetime
    wb.properties.creator = "E-Commerce Accounting (synthetic)"
    wb.properties.created = datetime(close_run.month.year, close_run.month.month, 1)
    path = out_dir / f"E-Commerce Allocation {ym}.xlsx"
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)

    key = f"""# ANSWER KEY - planted re-keying error in `E-Commerce Allocation {ym}.xlsx`

Synthetic manual allocation workbook (what the team builds by hand today, slide 39), generated from the automated
close run (rules {close_run.rules_version}, seed {seed}). It matches the automated journal line for line, **except one
orange input cell where two adjacent digits were transposed** while re-keying a source report.

| | |
|---|---|
| Document | `{err['doc_no']}` |
| Line | {err['description']} |
| Orange input cell | `{err['input_cell']}` |
| Journal Entries cell | `{err['je_cell']}` |
| Correct amount (source report) | {err['correct']:,.2f} |
| Keyed amount (workbook) | {err['keyed']:,.2f} |
| Difference (keyed - correct) | {err['delta']:,.2f} |

Because the workbook's formulas derive the balancing line from the inputs, the balancing line of the same document is
also off by the same amount and the document **still balances** (which is why nobody notices):

| | |
|---|---|
| Balancing line | {err['balancing_description']} (`{err['balancing_je_cell']}`) |
| Correct | {err['balancing_correct']:,.2f} |
| Workbook | {err['balancing_keyed']:,.2f} |

Expected result of the workbook comparison: every document and line matches except these two lines in
`{err['doc_no']}`; the root cause is the single orange cell `{err['input_cell']}` (digits transposed).
"""
    (out_dir / "ANSWER_KEY.md").write_text(key)
    return path


def _summary(run: CloseRun) -> str:
    out = [f"Close {run.month:%Y-%m}  rules {run.rules_version}  lines {len(run.allocations)}"]
    out.append(run.doc_totals().to_string(index=False))
    out.append("\nsource totals:")
    out.append(pd.DataFrame(run.source_totals).to_string(index=False))
    out.append("\nAR invoices (revenue) / customer balances:")
    out.append(run.customer_balances().to_string(index=False))
    out.append(f"\nexceptions ({len(run.exceptions)}):")
    for e in run.exceptions:
        out.append(f"  [{e['severity']}] {e['rule_id']}: {e['message']}")
    return "\n".join(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Run the e-commerce month-end close rules")
    ap.add_argument("--month", default="2026-09")
    ap.add_argument("--harmonized", default=str(DEFAULT_HARMONIZED))
    ap.add_argument("--finance", default=str(DEFAULT_FINANCE))
    ap.add_argument("--rules", default=str(DEFAULT_RULES))
    ap.add_argument("--workbook", action="store_true", help="also write the synthetic manual allocation workbook")
    ap.add_argument("--seed", type=int, default=7)
    a = ap.parse_args(argv)
    y, m = map(int, a.month.split("-"))
    run = run_close(date(y, m, 1), a.harmonized, a.finance, a.rules)
    print(_summary(run))
    if a.workbook:
        out_dir = DATA_DIR / "close_inputs" / f"{y:04d}-{m:02d}"
        path = write_manual_workbook(run, out_dir, seed=a.seed)
        print(f"\nmanual workbook: {path}")


if __name__ == "__main__":
    main()
