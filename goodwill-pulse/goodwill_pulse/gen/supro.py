"""Supro (store point-of-sale) daily sales per store, for the Revenue Hub's store side.

Real numbers come from the weekly Supro "Daily Sales Sheet" workbooks (demo_data/05_store_weekly_sales_supro/*.xlsx,
Data Sheet tab: net sales and customers per store per day). Those sheets have no returns or units, so both are estimated
on every day. Days with no workbook are synthesized: seeded per store x day (same output every run), scaled by each
store's typical daily sales (demo_pack.STORES, the same base the workbooks are generated from), weekday pattern
(Saturday busiest), month seasonality (December highest), and closed on Thanksgiving, Christmas Day and Easter.
"""
from __future__ import annotations

import random
from datetime import date, datetime, timedelta
from pathlib import Path

from .demo_pack import DOW_MULT, OUT as DEMO_DIR, STORES

COLS = ["date", "store", "sales", "customers", "returns", "units", "est"]
DEFINITIONS = {
    "sales": "Net in-store register sales for the day in dollars (after returns), as rung up in Supro.",
    "customers": "Register transactions that day (one per checkout), which Supro reports as customers.",
    "returns": "Dollar value of in-store returns refunded that day. Estimated on every day (about 0.5-1.5% of sales): "
               "the weekly Supro sheet does not report returns.",
    "units": "Items sold at the register that day. Estimated on every day (about 3-4 items per customer): "
             "the weekly Supro sheet does not report units.",
    "est": "1 when the day's sales and customers are estimated (no Supro report loaded for that day), "
           "0 when they are taken from a weekly Supro Daily Sales Sheet.",
    "dates": "Business dates in America/New_York. Stores are closed (all zeros) on Thanksgiving, Christmas Day and Easter.",
}
SUPRO_DIR = DEMO_DIR / "05_store_weekly_sales_supro"
BASE = {sid: base for sid, _, _, base in STORES}      # typical daily net sales per store
DEFAULT_BASE = 2000
SEASON = {1: 0.84, 2: 0.88, 3: 1.00, 4: 1.03, 5: 1.04, 6: 0.98, 7: 0.97, 8: 1.03, 9: 1.00, 10: 1.04, 11: 1.10, 12: 1.24}


def _easter(y: int) -> date:
    """Western Easter (anonymous Gregorian algorithm)."""
    a, b, c = y % 19, y // 100, y % 100
    d, e = b // 4, b % 4
    g = (8 * b + 13) // 25
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l_ = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l_) // 451
    mo, dy = divmod(h + l_ - 7 * m + 114, 31)
    return date(y, mo, dy + 1)


def _closed(d: date) -> bool:
    thanksgiving = date(d.year, 11, 1) + timedelta(days=(3 - date(d.year, 11, 1).weekday()) % 7 + 21)
    return d in (thanksgiving, date(d.year, 12, 25), _easter(d.year))


def _day_mult(d: date) -> float:
    m = DOW_MULT[(d.weekday() + 1) % 7] * SEASON[d.month]
    if d.month == 11 and d.weekday() == 4 and 23 <= d.day <= 29:      # Black Friday
        m *= 1.6
    elif d.month == 12 and 18 <= d.day <= 24:                         # last week before Christmas
        m *= 1.15
    return m


def read_reports(folder: Path | None = None) -> tuple[dict[tuple[str, str], tuple[float, int]], dict[str, str]]:
    """Net sales and customers by (date, store_id) from the weekly Supro workbooks, plus the file each date came from."""
    from openpyxl import load_workbook
    got, files = {}, {}
    folder = folder or SUPRO_DIR
    for p in sorted(folder.glob("Weekly Sales *.xlsx")) if folder.is_dir() else []:
        wb = load_workbook(p, read_only=True, data_only=True)
        if "Data Sheet" not in wb.sheetnames:
            continue
        it = wb["Data Sheet"].iter_rows(values_only=True)
        hdr = [str(h or "").strip() for h in next(it, [])]
        ix = {k: hdr.index(k) for k in ("Store ID", "Date", "Net Sales", "Customers") if k in hdr}
        if len(ix) < 4:
            continue
        for r in it:
            sid, d, s, c = (r[ix[k]] for k in ("Store ID", "Date", "Net Sales", "Customers"))
            if not sid or d is None or s is None:
                continue
            d = (d.date() if isinstance(d, datetime) else d).isoformat()
            got[(d, str(sid))] = (float(s), int(c or 0))
            files[d] = p.name
        wb.close()
    return got, files


def supro_days(con, first: date, last: date) -> tuple[list[list], dict[str, dict]]:
    """Every store x every day in [first, last], in COLS order, plus per-month info ({"YYYY-MM": {estimated, note, ...}})."""
    stores = [s for (s,) in con.execute("SELECT store_id FROM dim_store ORDER BY 1").fetchall()]
    real, files = read_reports(SUPRO_DIR)
    rows, months = [], {}
    d = first
    while d <= last:
        ds, m = d.isoformat(), d.strftime("%Y-%m")
        mi = months.setdefault(m, {"real": set(), "est": set(), "files": set()})
        for sid in stores:
            rng = random.Random(f"supro-{sid}-{ds}")
            if (ds, sid) in real:
                sales, cust = real[(ds, sid)]
                est = 0
                mi["real"].add(ds)
                mi["files"].add(files[ds])
            elif _closed(d):
                rows.append([ds, sid, 0.0, 0, 0.0, 0, 1])
                mi["est"].add(ds)
                continue
            else:
                sales = BASE.get(sid, DEFAULT_BASE) * _day_mult(d) * rng.uniform(0.82, 1.18)
                cust = max(1, round(sales / rng.uniform(17.5, 24.5)))
                est = 1
                mi["est"].add(ds)
            rows.append([ds, sid, round(sales, 2), int(cust), round(sales * rng.uniform(0.005, 0.015), 2),
                         int(round(cust * rng.uniform(3.0, 4.2))), est])
        d += timedelta(days=1)
    info = {}
    for m, mi in months.items():
        if not mi["real"]:
            note = "Estimated: no Supro report is loaded for this month, so sales and customers are modeled from each store's typical day."
        elif not mi["est"]:
            note = "Sales and customers come from the weekly Supro Daily Sales Sheets; returns and units are estimated."
        else:
            note = (f"Sales and customers for {min(mi['real'])} to {max(mi['real'])} come from the weekly Supro Daily Sales Sheets; "
                    "other days are estimated, and returns and units are estimated on every day.")
        info[m] = {"estimated": bool(mi["est"]), "note": note, "reported_days": len(mi["real"]),
                   "estimated_days": len(mi["est"]), "reports": sorted(mi["files"])}
    return rows, info
