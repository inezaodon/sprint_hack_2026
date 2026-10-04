"""Test reports for one new business day, in the same layouts as the Goodwill test reports, ready for the Upload tab.

    .venv/bin/python -m goodwill_pulse.gen.today_reports                 # today (America/New_York)
    .venv/bin/python -m goodwill_pulse.gen.today_reports --day 2026-10-04 --copy-to ~/Desktop/Goodwill_test_reports

Writes demo_data/07_today_<day>/:
  paid_orders_<day-1>_<day>.xlsx           Upright paid orders (Pacific times, totals row). The file name spans the day
                                           before because Upright days are Pacific: Pacific <day-1> 00:00 .. <day+1> 00:00
                                           is what fully covers the Eastern business day, so the Upload tab adds it.
  paid_orders_<day-1>_<day>_differs.xlsx   the same report with one subtotal changed and one row duplicated; upload it
                                           after the first to watch the day's ShopGoodwill numbers change
  cashmonkey_orders_<day>.xlsx             Cash Monkey orders, one row per unit, UTC times
  messy_sales_export_<day>.xlsx            a hand-made export: title rows, '$' text amounts, a test order, bad rows, Total row
  expected.json                            what each file should add for the day, by marketplace

The Upload tab only adds a business day the file fully covers (midnight to midnight Eastern), so the Cash Monkey and
messy files also carry one order just before and one just after the day; those edge days stay "partial" and are ignored.
Everything is synthetic.
"""
from __future__ import annotations

import argparse
import json
import random
import shutil
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font

from ..config import ROOT
from .world import World
from .writers import CASHMONKEY_HEADER, UPRIGHT_HEADER, upright_rows

ET, PT, UTC = ZoneInfo("America/New_York"), ZoneInfo("America/Los_Angeles"), ZoneInfo("UTC")
OUT_ROOT = ROOT / "demo_data"
LAST_UPRIGHT_ID, LAST_SGW_ORDER = 24_853_241, 66_176_736      # newest ids in the warehouse (Sat 10/03 report)
UPRIGHT_CHANNEL = {"Shopgoodwill": "shopgoodwill", "eBay": "ebay", "GoodwillFinds": "goodwillfinds"}
CM_CHANNEL = {"Amazon-MF": "amazon", "eBay": "ebay", "Goodwillbooks": "goodwillbooks"}


def _num(s):
    return float(s) if s not in ("", None) else None


def _et_day(dt: datetime) -> date:
    return dt.astimezone(ET).date()


def _book(path: Path, title: str, header: list[str], rows: list[list], date_col: int | None = None,
          date_fmt: str = "mm/dd/yyyy hh:mm:ss", lead: list[list] | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in lead or []:
        ws.append(r)
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
        if date_col is not None and isinstance(r[date_col], datetime):
            ws.cell(ws.max_row, date_col + 1).number_format = date_fmt
    wb.save(path)


def upright(world: World, day: date) -> tuple[list[list], list]:
    """Orders whose Eastern business day is `day`, as typed Upright rows (+ the Order objects)."""
    orders = [o for d in (day - timedelta(days=1), day) for o in world.upright_orders(d) if _et_day(o.paid_at) == day]
    for o in orders:                                   # ShopGoodwill buyer ids are numeric, like the real export
        if o.channel_raw == "Shopgoodwill":
            o.buyer = str(world.rng.randint(1_000_000, 2_999_999))
    rows = []
    for r in upright_rows(orders):
        rows.append([int(r[0]), r[1], r[2], None, r[4], int(r[5]), None, r[6 + 1], _num(r[8]), _num(r[9]), _num(r[10]),
                     None, _num(r[12]), _num(r[13]), 0, "USD", _num(r[16]), _num(r[17]),
                     datetime.strptime(r[18], "%m/%d/%Y %H:%M:%S")])
    return rows, orders


def _with_totals(rows: list[list]) -> list[list]:
    tot = [None] * len(UPRIGHT_HEADER)
    tot[9] = round(sum(r[9] for r in rows), 2)
    tot[10] = round(sum(r[10] or 0 for r in rows), 2)
    return rows + [tot]


def cashmonkey(world: World, day: date) -> list[list]:
    """Book orders from just before to just after the Eastern day (UTC times, one row per unit)."""
    start = datetime.combine(day, time(0), ET) - timedelta(minutes=30)
    end = datetime.combine(day + timedelta(days=1), time(0), ET) + timedelta(minutes=20)
    orders = [o for d in (day, day + timedelta(days=1)) for o in world.cashmonkey_orders(d) if start <= o.paid_at <= end]
    orders.sort(key=lambda o: o.paid_at)
    # guarantee one order on each side of the day's Eastern midnights, so the day counts as fully covered
    orders[0].paid_at = start + timedelta(minutes=11)
    orders[-1].paid_at = end - timedelta(minutes=7)
    rng = random.Random(day.toordinal())
    rows = []
    for o in orders:
        fee_rate = {"Amazon-MF": 0.15, "eBay": 0.1325, "Goodwillbooks": 0.05}[o.channel_raw]
        credit = round(o.shipping / o.item_count, 2)
        for line in o.lines:
            fees = round(line["price"] * fee_rate + (1.80 if o.channel_raw == "Amazon-MF" else 0.30), 2)
            cost = round(rng.uniform(3.2, 4.4), 2)
            rows.append([o.paid_at.astimezone(UTC).replace(tzinfo=None, microsecond=0), o.order_id, "276 - Goodwill Michiana",
                         o.channel_raw, line["sku"], rng.randint(100_000_000, 999_999_999), line["title"],
                         rng.choice(["Used - Good", "Used - Very Good", "Used - Acceptable", "Used - Like New"]),
                         line["store"], 1, line["price"], credit, fees, cost,
                         round(line["price"] + credit - fees - cost, 2), o.buyer])
    return rows


def messy(orders: list, day: date, rng: random.Random) -> tuple[list[list], list[list]]:
    """GoodwillFinds orders of the day (same ids and amounts as the Upright file) in a hand-made export layout."""
    gwf = sorted((o for o in orders if o.channel_raw == "GoodwillFinds"), key=lambda o: o.paid_at)

    def line(oid, at, sub, ship, fee, buyer):
        return [oid, "Goodwillfinds", at.astimezone(ET).strftime("%m/%d/%Y %H:%M"), f"${sub:,.2f}", ship, fee, buyer]

    rows = []
    edge0 = datetime.combine(day, time(0), ET) - timedelta(minutes=6)        # yesterday 23:54 ET (partial day)
    rows.append(line(str(LAST_SGW_ORDER + 9001), edge0, 48.50, 12.0, 1.71, "2210457"))
    for o in gwf:
        handling = 3.0 * o.item_count
        tax = round(o.subtotal * o.tax_rate, 2)
        total = round(o.subtotal + o.shipping + handling + tax, 2)
        pay_fee = round(total * 0.029 + 0.30, 2) if o.payment_type == "PayPal" else 0.0
        rows.append(line(o.order_id, o.paid_at, o.subtotal, round(o.shipping + handling, 2), pay_fee, o.buyer))
    edge1 = datetime.combine(day + timedelta(days=1), time(0), ET) + timedelta(minutes=4)   # tomorrow 00:04 ET
    rows.append(line(str(LAST_SGW_ORDER + 9002), edge1, 36.20, 9.5, 0.0, "1834402"))
    valid = list(rows)
    rows += [[None] * 7,
             ["TEST-1", "Goodwillfinds", f"{day:%m/%d/%Y} 10:00", 99, 5, 1, "TEST"],
             ["BAD-1", "Goodwillfinds", "not a date", 10, 1, 1, "x"],
             ["BAD-2", "Goodwillfinds", f"{day:%m/%d/%Y} 11:00", "ten dollars", 1, 1, "y"]]
    total = ["Total", None, None, round(sum(float(r[3].strip("$").replace(",", "")) for r in valid), 2),
             round(sum(r[4] for r in valid), 2), round(sum(r[5] for r in valid), 2), None]
    return rows + [total], gwf


def build(day: date, seed: int = 1004, copy_to: Path | None = None) -> Path:
    world = World(seed)
    world._upright_id, world._channel_id = LAST_UPRIGHT_ID + 5, LAST_SGW_ORDER + 3
    out = OUT_ROOT / f"07_today_{day:%Y-%m-%d}"
    shutil.rmtree(out, ignore_errors=True)
    out.mkdir(parents=True)
    prev = day - timedelta(days=1)
    stem = f"paid_orders_{prev:%m-%d-%Y}_{day:%m-%d-%Y}"
    expected = {"business_date": day.isoformat(), "files": {}}

    up_rows, up_orders = upright(world, day)
    _book(out / f"{stem}.xlsx", "Paid orders", UPRIGHT_HEADER, _with_totals(up_rows), date_col=18)
    agg = defaultdict(lambda: {"orders": 0, "item_sales": 0.0})
    for o in up_orders:
        a = agg[UPRIGHT_CHANNEL[o.channel_raw]]
        a["orders"] += 1
        a["item_sales"] = round(a["item_sales"] + o.subtotal, 2)
    expected["files"][f"{stem}.xlsx"] = dict(agg)

    rng = random.Random(seed)
    differs = [list(r) for r in up_rows]
    i = next(k for k, r in enumerate(differs) if r[1] == "Shopgoodwill" and r[9] > 40)
    differs[i][9] = round(differs[i][9] + 25.0, 2)             # a re-keyed subtotal
    differs.insert(rng.randrange(len(differs)), list(differs[rng.randrange(len(differs))]))   # a duplicated row
    _book(out / f"{stem}_differs.xlsx", "Paid orders", UPRIGHT_HEADER, _with_totals(differs), date_col=18)
    expected["files"][f"{stem}_differs.xlsx"] = {"change": f"order {differs[i][2]} subtotal +25.00; one row duplicated"}

    cm_rows = cashmonkey(world, day)
    _book(out / f"cashmonkey_orders_{day:%Y-%m-%d}.xlsx", "Orders", CASHMONKEY_HEADER, cm_rows, date_col=0,
          date_fmt="yyyy-mm-dd hh:mm:ss")
    cm = defaultdict(lambda: {"orders": set(), "item_sales": 0.0})
    for r in cm_rows:
        if _et_day(r[0].replace(tzinfo=UTC)) == day:
            a = cm[CM_CHANNEL[r[3]]]
            a["orders"].add(r[1])
            a["item_sales"] = round(a["item_sales"] + r[10], 2)
    expected["files"][f"cashmonkey_orders_{day:%Y-%m-%d}.xlsx"] = {
        k: {"orders": len(v["orders"]), "item_sales": v["item_sales"]} for k, v in cm.items()}

    m_rows, gwf = messy(up_orders, day, rng)
    _book(out / f"messy_sales_export_{day:%Y-%m-%d}.xlsx", "Export",
          ["Order #", "Marketplace", "Date Ordered", "Item Subtotal", "Shipping Charged", "Marketplace Fee", "Buyer"],
          m_rows, lead=[[f"Daily sales export - Goodwill Michiana - {day:%a %b %-d, %Y}"] + [None] * 6, [None] * 7])
    expected["files"][f"messy_sales_export_{day:%Y-%m-%d}.xlsx"] = {
        "goodwillfinds": {"orders": len(gwf), "item_sales": round(sum(o.subtotal for o in gwf), 2)},
        "note": "same GoodwillFinds orders as the Upright file; TEST-1, BAD-1, BAD-2 and the edge-day orders are not counted"}
    (out / "expected.json").write_text(json.dumps(expected, indent=1))

    if copy_to:
        copy_to.mkdir(parents=True, exist_ok=True)
        for p in out.glob("*.xlsx"):
            shutil.copy2(p, copy_to / p.name)
    return out


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--day", type=date.fromisoformat, default=datetime.now(ET).date())
    p.add_argument("--seed", type=int, default=1004)
    p.add_argument("--copy-to", type=Path)
    a = p.parse_args()
    out = build(a.day, a.seed, a.copy_to.expanduser() if a.copy_to else None)
    print(out)
    print((out / "expected.json").read_text())


if __name__ == "__main__":
    main()
