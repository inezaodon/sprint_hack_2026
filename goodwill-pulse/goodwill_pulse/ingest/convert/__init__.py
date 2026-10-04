"""Turn a file that isn't a spreadsheet into spreadsheets the existing pipeline can load.

    from goodwill_pulse.ingest.convert import convert
    conv = convert(Path("Goodwill Books statement.eml"))
    conv.write(out_dir)          # one .xlsx per table found; CSV / Excel attachments copied unchanged

    python -m goodwill_pulse.ingest.convert FILE [--out DIR]

Why: ingest/pipeline.py reads CSV and Excel only, but slide 38's sources also arrive as email deliveries (Upright),
emailed PDF statements (Goodwill Books), bank downloads (OFX), portal pages (HTML), Amazon flat files (tab-delimited
.txt) and zips. This module changes the FORMAT only: every value is copied as text, nothing is computed. Which report
a table is, and what its columns mean, is decided afterwards by the same recognizers as an uploaded spreadsheet.

How each format is read (deterministic first; Claude only when nothing else can read it):
  csv, xlsx, xls     passed through unchanged
  txt, tsv, psv      delimiter sniffed; title lines above the header kept as notes (Amazon date range report)
  html               every data <table> (layout tables in emails are skipped)
  json, xml          the list of records, one row each
  ofx, qfx           bank transactions, in the bank activity layout's column names
  eml (msg)          attachments converted recursively; tables in the body; sender, date, subject and Message-ID
                     kept for traceability. Outlook .msg needs the optional extract-msg package.
  zip                every file inside converted
  pdf                tables found under a known layout's column headers, or ruled tables; else Claude
  png, jpg, gif      Claude (a screenshot or phone photo of a report)
Anything Claude read is marked method="ai"; ingest/intake.py holds it for a person to confirm before it loads.

Each converted workbook: sheet 1 = the table (all cells text, so SKUs like '00123' survive), then 'Stated values'
(totals printed outside the table, when any), then 'About this file' (source, SHA-256, email headers, how it was read).
"""
from __future__ import annotations

import email
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from email import policy
from pathlib import Path

from ...config import finance_layouts, mappings
from . import ai_read, pdf as pdf_mod, tabular

PASSTHROUGH = {"csv", "xlsx", "xls"}
MAX_DEPTH = 4                      # email inside a zip inside an email ...
MIN_IMAGE_BYTES = 20_000           # smaller images in an email are logos / signatures, not reports
DATA_SHEET, STATED_SHEET, ABOUT_SHEET = "Data", "Stated values", "About this file"


class ConversionError(ValueError):
    """The file can't be converted; str(err) is safe to show a Goodwill user."""


@dataclass
class Output:
    name: str                                   # file name to write
    origin: str                                 # where in the upload it came from
    method: str                                 # passthrough | parsed | ai
    header: list[str] | None = None
    rows: list[list[str]] | None = None
    raw: bytes | None = None                    # passthrough bytes
    layout: str | None = None                   # layout id when the reader already knows it
    stated: dict[str, str] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    path: Path | None = None

    @property
    def row_count(self) -> int:
        return len(self.rows or [])


@dataclass
class Conversion:
    source_name: str
    sha256: str
    fmt: str
    envelope: dict[str, str] = field(default_factory=dict)    # email from / to / subject / date / message_id
    outputs: list[Output] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)            # what couldn't be read, and why

    @property
    def needs_review(self) -> bool:
        return any(o.method == "ai" for o in self.outputs)

    def write(self, out_dir: Path) -> list[Path]:
        out_dir = Path(out_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        for o in self.outputs:
            o.path = out_dir / o.name
            if o.raw is not None:
                o.path.write_bytes(o.raw)
            else:
                _write_xlsx(self, o, o.path)
        return [o.path for o in self.outputs]

    def summary(self) -> dict:
        return {"source": self.source_name, "sha256": self.sha256, "format": self.fmt, "envelope": self.envelope,
                "notes": self.notes,
                "outputs": [{"name": o.name, "origin": o.origin, "method": o.method, "layout": o.layout,
                             "rows": o.row_count if o.raw is None else None, "stated": o.stated, "notes": o.notes,
                             "path": str(o.path) if o.path else None} for o in self.outputs]}


# ------------------------------------------------------------------------------------------------------ layouts
def known_layouts() -> list[dict]:
    """Every layout the system can load (sales reports + month-end inputs): what the PDF reader anchors on and what
    Claude is told to use. [{id, label, kind, columns, required, aliases}]"""
    out = []
    for lid, spec in sorted(finance_layouts().items()):
        out.append({"id": lid, "label": spec["label"], "kind": "finance", "columns": list(spec["columns"]),
                    "required": [c for c, m in spec["columns"].items() if m.get("required")],
                    "aliases": dict(spec.get("aliases") or {})})
    for rt, spec in sorted(mappings().items()):
        out.append({"id": rt, "label": spec["label"], "kind": "sales", "columns": list(spec["columns"]),
                    "required": [c for c, m in spec["columns"].items() if m.get("required")],
                    "aliases": dict(spec.get("aliases") or {})})
    return out


# -------------------------------------------------------------------------------------------------------- sniff
def sniff(data: bytes, name: str) -> str:
    """Format from the bytes first (a .csv that is really Excel still reads), the extension second."""
    ext = Path(name).suffix.lower().lstrip(".")
    head = data[:8192]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            return "unknown"
        if "[Content_Types].xml" in names and any(n.startswith("xl/") for n in names):
            return "xlsx"
        if "[Content_Types].xml" in names and any(n.startswith("word/") for n in names):
            return "docx"
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "msg" if ext == "msg" else "xls"
    if head.startswith(b"\x89PNG"):
        return "png"
    if head.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if head[:6] in (b"GIF87a", b"GIF89a"):
        return "gif"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "webp"
    text = tabular.decode(head).lstrip()
    low = text[:2000].lower()
    if ext == "eml" or (re.match(r"[A-Za-z-]+:", text) and re.search(r"(?im)^from:", text)
                        and re.search(r"(?im)^(subject|date|message-id|mime-version):", text)):
        return "eml"
    if "ofxheader" in low[:300] or "<ofx>" in low:
        return "ofx"
    if ext in ("html", "htm") or re.search(r"<(html|table|body)[\s>]", low):
        return "html"
    if text.startswith("<?xml") or ext == "xml":
        return "xml"
    if text[:1] in "[{" or ext in ("json", "jsonl", "ndjson"):
        return "json"
    if ext == "csv":
        return "csv"
    if ext in ("txt", "tsv", "tab", "psv", "dat", "") or text:
        return "text"
    return "unknown"


# ------------------------------------------------------------------------------------------------------ convert
def convert(path: Path | str | None = None, *, data: bytes | None = None, name: str | None = None) -> Conversion:
    if data is None:
        path = Path(path)
        data = path.read_bytes()
    name = name or (Path(path).name if path else "upload")
    conv = Conversion(name, hashlib.sha256(data).hexdigest(), sniff(data, name))
    _Reader(conv).read(data, name, origin="", depth=0)
    _unique_names(conv.outputs)
    if not conv.outputs and not conv.notes:
        conv.notes.append(f"No table found in {name}.")
    return conv


def _stem(name: str) -> str:
    stem = Path(name).stem if Path(name).suffix else name
    stem = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", stem).strip(" .")
    return re.sub(r"\s+", " ", stem)[:120] or "converted"


def _unique_names(outputs: list[Output]) -> None:
    seen: dict[str, int] = {}
    for o in outputs:
        key = o.name.lower()
        if key in seen:
            seen[key] += 1
            p = Path(o.name)
            o.name = f"{p.stem} ({seen[key]}){p.suffix}"
        else:
            seen[key] = 1


def _parts(part):
    """Leaf parts of an email, in order. A forwarded email (message/rfc822) is yielded whole, not walked into, so
    its attachments are read once, by the recursive call, with their own envelope."""
    if part.get_content_type() == "message/rfc822" or not part.is_multipart():
        yield part
        return
    for p in part.iter_parts():
        yield from _parts(p)


class _Reader:
    def __init__(self, conv: Conversion):
        self.conv = conv
        self._layouts: list[dict] | None = None

    @property
    def layouts(self) -> list[dict]:
        if self._layouts is None:
            self._layouts = known_layouts()
        return self._layouts

    def _where(self, origin: str, name: str) -> str:
        return f"{origin} > {name}" if origin else name

    def table(self, name: str, origin: str, method: str, header: list[str], rows: list[list[str]],
              layout: str | None = None, stated: dict | None = None, notes: list[str] | None = None) -> None:
        self.conv.outputs.append(Output(f"{_stem(name)}.xlsx", origin, method, list(header), rows, layout=layout,
                                        stated=dict(stated or {}), notes=list(notes or [])))

    # --- dispatch ---------------------------------------------------------------------------------------------
    def read(self, data: bytes, name: str, origin: str, depth: int) -> None:
        kind = sniff(data, name)
        where = self._where(origin, name)
        if depth > MAX_DEPTH:
            self.conv.notes.append(f"{where}: nested too deep; skipped.")
            return
        handler = getattr(self, f"_{kind}", None)
        if kind in PASSTHROUGH:
            self.conv.outputs.append(Output(Path(name).name if Path(name).suffix else f"{_stem(name)}.{kind}",
                                            where, "passthrough", raw=data))
        elif handler is None:
            self.conv.notes.append(f"{where}: can't read this kind of file ({kind}). Save it as CSV, Excel or PDF.")
        else:
            handler(data, name, where, depth)

    # --- text formats -----------------------------------------------------------------------------------------
    def _text(self, data, name, where, depth):
        got = tabular.read_delimited(tabular.decode(data))
        if got is None:
            return self._ai_text(tabular.decode(data), name, where, "no columns found in the text")
        (header, rows), preamble = got
        self.table(name, where, "parsed", header, rows,
                   notes=[f"Text above the table: {' | '.join(preamble)}"] if preamble else None)

    def _html(self, data, name, where, depth):
        text = tabular.decode(data)
        tables = tabular.read_html(text)
        for i, (header, rows) in enumerate(tables, 1):
            self.table(name if len(tables) == 1 else f"{_stem(name)} (table {i})", f"{where}, table {i}", "parsed",
                       header, rows)
        if not tables:
            self._ai_text(tabular.html_text(text), name, where, "no data table in the page")

    def _json(self, data, name, where, depth):
        got = tabular.read_json(tabular.decode(data))
        if got is None:
            self.conv.notes.append(f"{where}: no list of records found in the JSON.")
            return
        self.table(name, where, "parsed", *got)

    def _xml(self, data, name, where, depth):
        got = tabular.read_xml(tabular.decode(data))
        if got is None:
            self.conv.notes.append(f"{where}: no repeated records found in the XML.")
            return
        self.table(name, where, "parsed", *got)

    def _ofx(self, data, name, where, depth):
        got = tabular.read_ofx(tabular.decode(data))
        if got is None:
            self.conv.notes.append(f"{where}: no transactions found in the bank download.")
            return
        self.table(name, where, "parsed", *got, layout="bank_0101")

    # --- containers -------------------------------------------------------------------------------------------
    def _zip(self, data, name, where, depth):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                base = Path(info.filename).name
                if info.is_dir() or info.filename.startswith("__MACOSX") or base.startswith("."):
                    continue
                self.read(z.read(info), base, where, depth + 1)

    def _eml(self, data, name, where, depth):
        msg = email.message_from_bytes(data, policy=policy.default)
        self._message(msg, name, where, depth)

    def _msg(self, data, name, where, depth):
        try:
            import extract_msg   # optional: pip install extract-msg
        except ImportError:
            self.conv.notes.append(f"{where}: Outlook .msg files need the extract-msg package. In Outlook, use "
                                   "File > Save As > .eml, or upload the attachment itself.")
            return
        m = extract_msg.Message(io.BytesIO(data))
        env = {"from": m.sender or "", "to": m.to or "", "subject": m.subject or "", "date": str(m.date or ""),
               "message_id": (m.messageId or "").strip()}
        self._envelope(env)
        found = len(self.conv.outputs)
        for a in m.attachments:
            fname = getattr(a, "longFilename", None) or getattr(a, "shortFilename", None) or "attachment"
            if isinstance(a.data, bytes):
                self.read(a.data, fname, f"{where} > attachment", depth + 1)
        self._body(m.htmlBody.decode("utf-8", "replace") if isinstance(m.htmlBody, bytes) else m.htmlBody,
                   m.body, env["subject"] or name, where, found)

    def _envelope(self, env: dict) -> None:
        if not self.conv.envelope:          # the outermost email is the one that arrived
            self.conv.envelope = {k: v for k, v in env.items() if v}

    def _message(self, msg, name, where, depth):
        env = {"from": str(msg.get("From", "")), "to": str(msg.get("To", "")), "subject": str(msg.get("Subject", "")),
               "date": str(msg.get("Date", "")), "message_id": str(msg.get("Message-ID", "")).strip()}
        self._envelope(env)
        found = len(self.conv.outputs)
        html_body = text_body = None
        for part in _parts(msg):
            ctype = part.get_content_type()
            if ctype == "message/rfc822":
                inner = part.get_payload(0) if part.is_multipart() else None
                if inner is not None:
                    self._message(inner, str(inner.get("Subject", "forwarded email")), f"{where} > forwarded", depth + 1)
                continue
            fname = part.get_filename()
            disp = part.get_content_disposition()
            if fname or disp == "attachment":
                payload = part.get_payload(decode=True) or b""
                if part.get_content_maintype() == "image" and (disp != "attachment" or len(payload) < MIN_IMAGE_BYTES):
                    continue                                    # logo / signature image
                self.read(payload, fname or "attachment", f"{where} > attachment", depth + 1)
            elif ctype == "text/html" and html_body is None:
                html_body = part.get_content()
            elif ctype == "text/plain" and text_body is None:
                text_body = part.get_content()
        self._body(html_body, text_body, env["subject"] or name, where, found)

    def _body(self, html_body, text_body, subject, where, found_before):
        """Tables in the email body. If the email had no attachment and no table, Claude reads the text."""
        tables = tabular.read_html(html_body) if html_body else []
        for i, (header, rows) in enumerate(tables, 1):
            self.table(f"{_stem(subject)} (email table {i})", f"{where}, email body table {i}", "parsed", header, rows)
        if tables or len(self.conv.outputs) > found_before:
            return
        text = text_body or (tabular.html_text(html_body) if html_body else "")
        got = tabular.read_delimited(text, delimiters=("\t", "|"), min_rows=2) if text else None
        if got:
            (header, rows), _ = got
            self.table(f"{_stem(subject)} (email text)", f"{where}, email body", "parsed", header, rows)
        elif text and text.strip():
            self._ai_text(text, subject, f"{where}, email body", "the email has no attachment or table")
        else:
            self.conv.notes.append(f"{where}: the email has no attachment and no table.")

    # --- PDF and images ---------------------------------------------------------------------------------------
    def _pdf(self, data, name, where, depth):
        try:
            res = pdf_mod.read_pdf(data, self.layouts)
        except Exception as e:  # noqa: BLE001  (damaged / encrypted PDFs)
            self.conv.notes.append(f"{where}: the PDF couldn't be opened ({type(e).__name__}).")
            return
        labels = {l["id"]: l["label"] for l in self.layouts}
        for t in res.tables:
            pages = f"page {t['pages'][0]}" if len(t["pages"]) == 1 else f"pages {t['pages'][0]}-{t['pages'][-1]}"
            suffix = f" ({labels[t['layout']]})" if t.get("layout") and len(res.tables) > 1 else ""
            self.table(f"{_stem(name)}{suffix}", f"{where}, {pages}", "parsed", t["header"], t["rows"],
                       layout=t.get("layout"), stated=res.stated)
        if not res.tables:
            why = "scanned PDF (no text layer)" if not res.text.strip() else "no table found in the PDF text"
            self._ai(name, where, why, pdf=data)

    def _image(self, data, name, where, depth):
        self._ai(name, where, "image of a report", image=(data, sniff(data, name)))

    _png = _jpeg = _gif = _webp = _image

    # --- Claude -----------------------------------------------------------------------------------------------
    def _ai_text(self, text, name, where, why):
        if not text.strip():
            self.conv.notes.append(f"{where}: {why}; nothing to read.")
            return
        self._ai(name, where, why, text=text)

    def _ai(self, name, where, why, **content):
        if not ai_read.available():
            self.conv.notes.append(f"{where}: {why}. Claude would read it, but AI is not configured "
                                   "(set ANTHROPIC_API_KEY), so nothing was extracted.")
            return
        try:
            got = ai_read.read(self.layouts, hint=f"{where} ({why})", **content)
        except ai_read.ai.AIUnavailable as e:
            self.conv.notes.append(f"{where}: {why}; Claude couldn't read it ({e}).")
            return
        for i, t in enumerate(got["tables"], 1):
            label = t["title"] or (f"{_stem(name)} (table {i})" if len(got["tables"]) > 1 else name)
            self.table(label, f"{where} (read by Claude)", "ai", t["header"], t["rows"], layout=t["layout"],
                       stated=got["stated"], notes=[got["notes"]] if got["notes"] else None)
        if not got["tables"]:
            self.conv.notes.append(f"{where}: Claude found no table. {got['notes']}".strip())


# ------------------------------------------------------------------------------------------------------- output
def _sheet_title(s: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", " ", s)[:31]


def _write_xlsx(conv: Conversion, o: Output, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    ws.append(o.header or [])
    for r in o.rows or []:
        ws.append(["" if c is None else str(c) for c in r])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, h in enumerate(o.header or [], 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in (o.rows or [])[:200] if i - 1 < len(r)])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(60, max(8, width + 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"                     # text: what you see is exactly what was in the source

    if o.stated:
        st = wb.create_sheet(_sheet_title(STATED_SHEET))
        st.append(["Label", "Value"])
        for k, v in o.stated.items():
            st.append([k, v])
        st.column_dimensions["A"].width, st.column_dimensions["B"].width = 32, 32

    about = wb.create_sheet(_sheet_title(ABOUT_SHEET))
    facts = [("Source file", conv.source_name), ("Source SHA-256", conv.sha256), ("Source format", conv.fmt),
             ("Found at", o.origin), ("Read by", {"parsed": "Deterministic parser (no AI)",
                                                  "ai": "Claude (transcribed; needs a person to confirm)"}.get(o.method, o.method)),
             ("Layout", o.layout or "not recognized yet"), ("Rows", str(o.row_count)),
             ("Converted at (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"))]
    facts += [(f"Email {k.replace('_', ' ')}", v) for k, v in conv.envelope.items()]
    facts += [("Note", n) for n in o.notes + conv.notes]
    for k, v in facts:
        about.append([k, v])
    about.column_dimensions["A"].width, about.column_dimensions["B"].width = 22, 100
    wb.save(path)
